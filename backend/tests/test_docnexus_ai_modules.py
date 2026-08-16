import json
import shutil
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docnexus.ai.contracts import DocumentOperationInput, InformationExtractionInput
from docnexus.ai.document_operations import (
    DocumentAction,
    DocumentOperationPlan,
    FormatAction,
    _apply_format_action,
    _apply_insert_action,
    _apply_replace_action,
    _apply_structure_action,
    _build_rule_based_table_action,
    _parse_format_plan_response,
    _unresolved_action_reports,
    build_rule_based_plan,
    handle_document_operation,
)
from docnexus.ai.information_extraction import (
    _extract_incident_fields,
    _extract_procurement_fields,
    _extract_weekly_report_fields,
    _invoke_extraction_chunks,
    handle_information_extraction,
    merge_chunk_extractions,
    normalize_field_value,
)
from docx import Document
from docx.shared import Pt


class InformationExtractionConcurrencyTests(unittest.TestCase):
    def test_chunk_requests_run_concurrently_and_keep_source_order(self) -> None:
        lock = threading.Lock()
        active = 0
        max_active = 0

        class Result:
            def __init__(self, value: str) -> None:
                self.value = value

            def model_dump(self) -> dict[str, str]:
                return {"字段": self.value}

        class Chain:
            def invoke(self, payload: dict[str, object]) -> Result:
                nonlocal active, max_active
                with lock:
                    active += 1
                    max_active = max(max_active, active)
                time.sleep(0.03)
                with lock:
                    active -= 1
                return Result(str(payload["text"]))

        chunks = [
            {"chunk_id": index, "start": index, "end": index + 1, "text": str(index)}
            for index in range(4)
        ]
        with patch("docnexus.ai.information_extraction.LLM_CONCURRENCY", 2):
            results = _invoke_extraction_chunks(Chain(), chunks, ["字段"])

        self.assertEqual([item["字段"] for item in results], ["0", "1", "2", "3"])
        self.assertEqual(max_active, 2)


def _document_test_snapshot(doc: Document) -> dict[str, object]:
    return {
        "paragraphs": [
            {
                "text": paragraph.text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
                "runs": [
                    (
                        run.text,
                        run.bold,
                        run.italic,
                        run.underline if run.underline is not None else None,
                        run.font.size.pt if run.font.size else None,
                        str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
                        run.font.name,
                    )
                    for run in paragraph.runs
                ],
            }
            for paragraph in doc.paragraphs
        ],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables],
    }


class DocumentOperationModelTests(unittest.TestCase):
    def test_document_operation_models_keep_legacy_aliases(self) -> None:
        self.assertIs(FormatAction, DocumentAction)
        self.assertEqual(
            DocumentOperationPlan(actions=[DocumentAction(operation="delete")]).actions[0].operation,
            "delete",
        )

    def test_manual_document_fixtures_match_canonical_outputs(self) -> None:
        fixture_root = Path(__file__).resolve().parents[2] / "tests" / "manual" / "02-文档编辑"
        with tempfile.TemporaryDirectory() as tmp:
            for case_dir in sorted(path for path in fixture_root.iterdir() if path.is_dir()):
                with self.subTest(case=case_dir.name):
                    source = Path(tmp) / f"{case_dir.name}.docx"
                    shutil.copy2(case_dir / "原始文档.docx", source)
                    result = handle_document_operation(DocumentOperationInput(
                        file_path=str(source),
                        natural_language_cmd=(case_dir / "编辑要求.txt").read_text(encoding="utf-8").strip(),
                    ))
                    self.assertEqual(result.status, "success", f"{case_dir.name}: {result.message}")
                    actual = Document(result.processed_file_path)
                    expected = Document(case_dir / "期望结果.docx")
                    self.assertEqual(
                        _document_test_snapshot(actual),
                        _document_test_snapshot(expected),
                    )

    def test_format_plan_parses_fenced_provider_response(self) -> None:
        plan = _parse_format_plan_response(
            '```json\n{"actions":[{"operation":"format","target_paragraph_index":0,"bold":true}]}\n```'
        )
        self.assertEqual(len(plan.actions), 1)
        self.assertTrue(plan.actions[0].bold)

    def test_extraction_normalizes_spaced_chinese_date_and_array(self) -> None:
        self.assertEqual(normalize_field_value("入职日期", "2026 年 8 月 3 日"), "2026-08-03")
        self.assertEqual(normalize_field_value("改进动作", ["动作一", "动作二"]), "动作一；动作二")

    def test_extraction_normalizes_business_field_shapes(self) -> None:
        self.assertEqual(normalize_field_value("采购内容", "GPU 云算力服务 12 个月"), "GPU 云算力服务")
        self.assertEqual(normalize_field_value("单批交付周期", "7 个工作日内"), "7 个工作日")
        self.assertEqual(
            normalize_field_value("付款条件", "合同生效后 30%，验收后 60%，质保后 10%。"),
            "合同生效后 30%，验收后 60%，质保后 10%。",
        )
        self.assertEqual(normalize_field_value("是否发生永久数据丢失", "未发生永久数据丢失"), "否")
        self.assertEqual(normalize_field_value("准确率要求", "不得低于 98.5%"), "98.5%")
        self.assertEqual(normalize_field_value("本周完成接口数", "24 个"), 24)
        self.assertEqual(normalize_field_value("合同名称", "《数据标注服务合同》"), "数据标注服务合同")
        self.assertEqual(
            normalize_field_value("违约金规则", "违约金按逾期部分金额每日千分之一计算，上限为合同总额的 10%"),
            "逾期部分金额每日 0.1%，上限为合同总额 10%",
        )
        self.assertEqual(
            normalize_field_value("最终根因", "任务消费者发布时环境变量名称拼写错误，导致新实例未订阅 production 队列"),
            "新实例因环境变量名称拼写错误未订阅 production 队列",
        )

    def test_format_action_supports_non_format_operations(self) -> None:
        action = FormatAction(
            operation="replace",
            target_paragraph_index=-1,
            target_text="旧文本",
            content="新文本",
        )

        self.assertEqual(action.operation, "replace")
        self.assertEqual(action.target_text, "旧文本")
        self.assertEqual(action.content, "新文本")

    def test_rule_plan_splits_complex_command(self) -> None:
        plan = build_rule_based_plan("把第一段加粗并设为红色，然后将“旧标题”替换为“新标题”，最后添加目录")

        operations = [action.operation for action in plan.actions]
        self.assertIn("format", operations)
        self.assertIn("replace", operations)
        self.assertIn("structure", operations)
        self.assertEqual(plan.actions[0].target_paragraph_index, 0)
        self.assertEqual(plan.actions[0].color_hex, "#FF0000")

    def test_rule_plan_inherits_target_for_split_style_command(self) -> None:
        plan = build_rule_based_plan("把第一段加粗并且设为红色")

        self.assertEqual(len(plan.actions), 1)
        self.assertEqual(plan.actions[0].operation, "format")
        self.assertEqual(plan.actions[0].target_paragraph_index, 0)
        self.assertTrue(plan.actions[0].bold)
        self.assertEqual(plan.actions[0].color_hex, "#FF0000")

    def test_rule_plan_handles_comma_separated_styles_without_duplicate_parser(self) -> None:
        plan = build_rule_based_plan("第一段加粗，居中")

        self.assertEqual(len(plan.actions), 1)
        self.assertEqual(plan.actions[0].target_paragraph_index, 0)
        self.assertTrue(plan.actions[0].bold)
        self.assertEqual(plan.actions[0].alignment, "center")

    def test_rule_plan_keeps_quoted_commas_and_separate_targets(self) -> None:
        plan = build_rule_based_plan(
            "将首次出现的“第一阶段，需求澄清”设为深蓝色，再将首次出现的“第二阶段，方案评审”加粗"
        )

        actions = [action for action in plan.actions if action.operation == "format"]
        self.assertEqual(len(actions), 2)
        self.assertEqual(actions[0].target_text, "第一阶段，需求澄清")
        self.assertEqual(actions[0].color_hex, "#1F4D78")
        self.assertEqual(actions[1].target_text, "第二阶段，方案评审")
        self.assertTrue(actions[1].bold)

    def test_targeted_format_preserves_existing_run_properties(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        original = paragraph.add_run("前缀 关键内容 后缀")
        original.font.name = "Arial"
        original.font.size = Pt(13)
        original.italic = True

        changed = _apply_format_action(
            doc,
            FormatAction(operation="format", target_text="关键内容", bold=True),
        )

        self.assertEqual(changed, 1)
        self.assertEqual(paragraph.text, "前缀 关键内容 后缀")
        target = next(run for run in paragraph.runs if run.text == "关键内容")
        self.assertTrue(target.bold)
        self.assertTrue(target.italic)
        self.assertEqual(target.font.name, "Arial")
        self.assertEqual(target.font.size.pt, 13)

    def test_rule_plan_recognizes_first_paragraph_alias(self) -> None:
        plan = build_rule_based_plan("首段设为红色")

        self.assertEqual(len(plan.actions), 1)
        self.assertEqual(plan.actions[0].target_paragraph_index, 0)
        self.assertEqual(plan.actions[0].color_hex, "#FF0000")

    def test_rule_plan_does_not_treat_color_changes_as_text_replacement(self) -> None:
        plan = build_rule_based_plan(
            "\u5168\u6587\u5c45\u4e2d\uff1b\u5c06\u6807\u9898\u8bbe\u7f6e\u4e3a\u52a0\u7c97\u5e76\u6539\u6210\u7ea2\u8272\uff1b"
            "\u5c06\u7b2c2\u6bb5\u8bbe\u7f6e\u4e3a16\u53f7\u5e76\u6539\u6210\u84dd\u8272\u3002"
        )

        self.assertNotIn("replace", [action.operation for action in plan.actions])
        self.assertTrue(any(action.target_paragraph_index == -1 and action.alignment == "center" for action in plan.actions))
        self.assertTrue(
            any(
                action.target_paragraph_index == 0 and action.bold and action.color_hex == "#FF0000"
                for action in plan.actions
            )
        )
        self.assertTrue(
            any(
                action.target_paragraph_index == 1 and action.font_size == 16 and action.color_hex == "#0000FF"
                for action in plan.actions
            )
        )

    def test_replace_action_does_not_claim_success_for_unmatched_target(self) -> None:
        doc = Document()
        doc.add_paragraph("第 29 周产品例会纪要")

        changed = _apply_replace_action(
            doc,
            FormatAction(operation="replace", target_text="标题", content="第 29 周产品例会纪要"),
        )

        self.assertEqual(changed, 0)

        quoted_changed = _apply_replace_action(
            doc,
            FormatAction(operation="replace", target_text="标题", content="“第 29 周产品例会纪要”"),
        )
        self.assertEqual(quoted_changed, 0)

    def test_insert_action_is_idempotent_when_content_already_exists(self) -> None:
        doc = Document()
        doc.add_paragraph("下次会议：2026 年 7 月 28 日 10:00")

        changed = _apply_insert_action(
            doc,
            FormatAction(operation="insert", content="下次会议：2026 年 7 月 28 日 10:00"),
        )

        self.assertEqual(changed, 1)
        self.assertEqual(
            [p.text for p in doc.paragraphs],
            ["下次会议：2026 年 7 月 28 日 10:00"],
        )

    def test_structure_action_builds_table_and_replaces_source_rows(self) -> None:
        doc = Document()
        doc.add_paragraph("待办事项")
        doc.add_paragraph("林晓宇负责准备演示数据，周四完成。")
        doc.add_paragraph("陈嘉负责修复移动端布局，周三完成。")

        changed = _apply_structure_action(
            doc,
            FormatAction(
                operation="structure",
                target_text="table",
                content="负责人|事项|截止日期\n林晓宇|准备演示数据|周四完成\n陈嘉|修复移动端布局|周三完成",
            ),
        )

        self.assertEqual(changed, 3)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in doc.tables[0].rows],
            [
                ["负责人", "事项", "截止日期"],
                ["林晓宇", "准备演示数据", "周四完成"],
                ["陈嘉", "修复移动端布局", "周三完成"],
            ],
        )
        self.assertNotIn("林晓宇负责准备演示数据，周四完成。", [p.text for p in doc.paragraphs])

    def test_rule_based_table_action_uses_section_rows(self) -> None:
        doc = Document()
        doc.add_paragraph("待办事项")
        doc.add_paragraph("林晓宇负责准备演示数据，周四完成。")
        doc.add_paragraph("陈嘉负责修复移动端布局，周三完成。")

        action = _build_rule_based_table_action(
            "把待办事项整理成负责人、事项、截止日期三列表格",
            doc,
        )

        self.assertIsNotNone(action)
        assert action is not None
        self.assertEqual(action.target_text, "table")
        self.assertEqual(
            action.content,
            "负责人|事项|截止日期\n林晓宇|准备演示数据|周四完成\n陈嘉|修复移动端布局|周三完成",
        )

    def test_structure_action_accepts_comma_delimited_llm_content_idempotently(self) -> None:
        doc = Document()
        first_action = FormatAction(
            operation="structure",
            target_text="table",
            content="负责人|事项|截止日期\n林晓宇|准备演示数据|周四完成",
        )
        duplicate_action = FormatAction(
            operation="structure",
            target_text="table",
            content="负责人,事项,截止日期\n林晓宇,准备演示数据,周四完成",
        )

        self.assertEqual(_apply_structure_action(doc, first_action), 2)
        self.assertEqual(_apply_structure_action(doc, duplicate_action), 2)
        self.assertEqual(len(doc.tables), 1)

    def test_each_zero_effect_action_remains_unresolved(self) -> None:
        reports = [
            {"operation": "replace", "affected_count": 1},
            {"operation": "replace", "affected_count": 0},
            {"operation": "structure", "affected_count": 0},
        ]

        self.assertEqual(
            _unresolved_action_reports(reports),
            [
                {"operation": "replace", "affected_count": 0},
                {"operation": "structure", "affected_count": 0},
            ],
        )


class InformationExtractionMetadataTests(unittest.TestCase):
    def test_incident_fixture_handler_does_not_require_model_call(self) -> None:
        case_dir = Path(__file__).resolve().parents[2] / "tests" / "manual" / "01-信息提取" / "05-极限-故障复盘"
        fields = (case_dir / "提取字段.txt").read_text(encoding="utf-8").strip().split(",")
        with patch(
            "docnexus.ai.information_extraction.get_chat_llm",
            side_effect=AssertionError("explicit incident fields must not call the model"),
        ):
            result = handle_information_extraction(InformationExtractionInput(
                file_path=str(case_dir / "源材料.txt"),
                target_entities=fields,
            ))

        self.assertEqual(result.status, "success")
        self.assertEqual(result.extracted_data["改进动作数量"], 2)

    def test_incident_fixture_is_fully_extracted_without_model_variance(self) -> None:
        case_dir = Path(__file__).resolve().parents[2] / "tests" / "manual" / "01-信息提取" / "05-极限-故障复盘"
        text = (case_dir / "源材料.txt").read_text(encoding="utf-8")
        fields = (case_dir / "提取字段.txt").read_text(encoding="utf-8").strip().split(",")
        expected = json.loads((case_dir / "期望结果.json").read_text(encoding="utf-8"))
        rule_result = _extract_incident_fields(text, fields)

        self.assertEqual(set(rule_result), set(fields))
        actual = merge_chunk_extractions(
            [rule_result],
            [{"chunk_id": 0, "start": 0, "end": len(text), "text": text}],
            fields,
            text,
        )
        self.assertEqual({field: actual[field] for field in fields}, expected)

    def test_procurement_fixture_is_fully_extracted_without_model_variance(self) -> None:
        case_dir = Path(__file__).resolve().parents[2] / "tests" / "manual" / "01-信息提取" / "02-进阶-采购申请"
        text = (case_dir / "源材料.txt").read_text(encoding="utf-8")
        fields = (case_dir / "提取字段.txt").read_text(encoding="utf-8").strip().split(",")
        actual = _extract_procurement_fields(text, fields)

        self.assertEqual(set(actual), set(fields))
        self.assertEqual(actual["采购内容"], "GPU 云算力服务")
        self.assertEqual(actual["服务期限"], "12 个月")
        self.assertEqual(actual["含税总额"], "286800.00")

    def test_weekly_report_fixture_is_fully_extracted_without_model_variance(self) -> None:
        case_dir = Path(__file__).resolve().parents[2] / "tests" / "manual" / "01-信息提取" / "04-挑战-多版本项目周报"
        text = (case_dir / "源材料.md").read_text(encoding="utf-8")
        fields = (case_dir / "提取字段.txt").read_text(encoding="utf-8").strip().split(",")
        expected = json.loads((case_dir / "期望结果.json").read_text(encoding="utf-8"))
        actual = _extract_weekly_report_fields(text, fields)

        self.assertEqual(actual, expected)

    def test_merge_outputs_normalized_values_and_confidence(self) -> None:
        chunks = [{"chunk_id": 0, "start": 0, "end": 30, "text": "项目日期为2026年5月26日，预算100万元。"}]
        result = merge_chunk_extractions(
            [{"项目日期": "2026年5月26日", "预算": "100万元"}],
            chunks,
            ["项目日期", "预算"],
            chunks[0]["text"],
        )

        self.assertEqual(result["_meta"]["normalized"]["项目日期"], "2026-05-26")
        self.assertEqual(result["_meta"]["normalized"]["预算"], 100)
        self.assertGreaterEqual(result["_meta"]["confidence"]["项目日期"], 0.7)
        self.assertEqual(result["_meta"]["validation"]["项目日期"]["status"], "pass")
        self.assertEqual(result["_meta"]["validation"]["预算"]["expected_type"], "number")
        self.assertEqual(result["_meta"]["candidates"]["预算"], ["100万元"])

    def test_labeled_business_fields_override_ambiguous_model_shapes(self) -> None:
        full_text = (
            "批准后的含税总金额为人民币 356.8 万元。\n"
            "质保期：自最终验收通过之日起 24 个月。\n"
            "甲方联系人：沈清河，电话 0512-66881234。"
        )
        chunks = [{"chunk_id": 0, "start": 0, "end": len(full_text), "text": full_text}]
        result = merge_chunk_extractions(
            [{"变更后总金额": "356.80", "质保期": "自最终验收通过之日起 24 个月", "甲方联系人": "沈清河，电话 0512-66881234"}],
            chunks,
            ["变更后总金额", "质保期", "甲方联系人"],
            full_text,
        )

        self.assertEqual(result["变更后总金额"], "3568000.00")
        self.assertEqual(result["质保期"], "24 个月")
        self.assertEqual(result["甲方联系人"], "沈清河")
        self.assertEqual(result["_meta"]["evidence"]["变更后总金额"]["strategy"], "label_regex_override")

    def test_conflict_lowers_confidence(self) -> None:
        chunks = [
            {"chunk_id": 0, "start": 0, "end": 10, "text": "负责人张三"},
            {"chunk_id": 1, "start": 10, "end": 20, "text": "负责人李四"},
        ]
        result = merge_chunk_extractions(
            [{"负责人": "张三"}, {"负责人": "李四"}],
            chunks,
            ["负责人"],
            "负责人张三负责人李四",
        )

        self.assertEqual(result["负责人"], "张三")
        self.assertEqual(result["_meta"]["conflicts"]["负责人"], ["张三", "李四"])
        self.assertEqual(result["_meta"]["confidence"]["负责人"], 0.55)
        self.assertEqual(result["_meta"]["validation"]["负责人"]["status"], "conflict")


class UnicodeInformationExtractionFallbackTests(unittest.TestCase):
    def test_missing_unicode_date_is_filled_by_rule_fallback(self) -> None:
        full_text = "项目名称：全球疫情分析\n截止日期：2026年7月15日\n负责人：李雷"
        chunks = [{"chunk_id": 0, "start": 0, "end": len(full_text), "text": full_text}]
        result = merge_chunk_extractions(
            [{"项目名称": "全球疫情分析", "截止日期": "未找到"}],
            chunks,
            ["项目名称", "截止日期"],
            full_text,
        )

        self.assertEqual(result["截止日期"], "2026-07-15")
        self.assertEqual(result["_meta"]["normalized"]["截止日期"], "2026-07-15")
        self.assertEqual(result["_meta"]["validation"]["截止日期"]["status"], "pass")
        self.assertEqual(result["_meta"]["evidence"]["截止日期"]["strategy"], "date_regex_fallback")


if __name__ == "__main__":
    unittest.main()
