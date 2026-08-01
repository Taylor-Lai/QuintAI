from __future__ import annotations

import json
import shutil
import tempfile
from pathlib import Path

import pytest
from docnexus.ai.document_operations import (
    _build_document_aware_plan,
    _execute_action,
    build_rule_based_plan,
    merge_rule_plans,
)
from docnexus.ai.table_engine.analyzers import DefaultTemplateAnalyzer
from docnexus.ai.table_engine.app import build_orchestrator
from docnexus.ai.table_engine.cli import discover_assets
from docnexus.ai.table_engine.config import AppConfig
from docnexus.ai.table_engine.core.models import EvidencePack, FileAsset
from docnexus.ai.table_engine.parsers import TextParser, XlsxParser
from docnexus.ai.table_engine.relational_records import build_template_anchored_records
from docnexus.ai.table_engine.writers import XlsxWriter
from docnexus.services.input_parsing import parse_user_list
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MANUAL_TABLES = ROOT / "tests" / "manual" / "03-表格填充"
MANUAL_DOCUMENTS = ROOT / "tests" / "manual" / "02-文档编辑"
MANUAL_EXTRACTIONS = ROOT / "tests" / "manual" / "01-信息提取"


def _workbook_snapshot(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False)
    result = {}
    for sheet in workbook.worksheets:
        cells = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                font_color = cell.font.color
                cells[cell.coordinate] = {
                    "value": cell.value,
                    "number_format": cell.number_format,
                    "font": (
                        cell.font.bold,
                        cell.font.italic,
                        str(font_color.rgb) if font_color is not None and font_color.type == "rgb" else None,
                    ),
                    "fill": str(cell.fill.fgColor.rgb),
                    "alignment": (cell.alignment.horizontal, cell.alignment.vertical),
                }
        result[sheet.title] = {
            "cells": cells,
            "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
        }
    return result


def _compare_workbook(actual: Path, expected: Path) -> list[dict[str, object]]:
    actual_snapshot = _workbook_snapshot(actual)
    expected_snapshot = _workbook_snapshot(expected)
    mismatches = []
    for sheet_name in sorted(set(actual_snapshot) | set(expected_snapshot)):
        if actual_snapshot.get(sheet_name) != expected_snapshot.get(sheet_name):
            mismatches.append({
                "sheet": sheet_name,
                "expected": expected_snapshot.get(sheet_name),
                "actual": actual_snapshot.get(sheet_name),
            })
    return mismatches


def _asset(path: Path, role: str) -> FileAsset:
    return FileAsset(
        id=path.name,
        path=str(path),
        name=path.name,
        ext=path.suffix.lstrip("."),
        role=role,
        mime_type=None,
        size=path.stat().st_size,
    )


def _build_records(case_name: str, work: Path):
    case_dir = MANUAL_TABLES / case_name
    work.mkdir(parents=True, exist_ok=True)
    template_path = work / "template.xlsx"
    shutil.copy2(case_dir / "空白模板.xlsx", template_path)
    template_doc = XlsxParser().parse(_asset(template_path, "template"))
    source_docs = []
    for source_path in [*case_dir.glob("源数据*"), *case_dir.glob("补充说明.txt")]:
        copied = work / source_path.name
        shutil.copy2(source_path, copied)
        parser = XlsxParser() if copied.suffix == ".xlsx" else TextParser()
        source_docs.append(parser.parse(_asset(copied, "source")))
    template_spec = DefaultTemplateAnalyzer().analyze(template_doc)
    records = build_template_anchored_records(
        template_doc,
        source_docs,
        template_spec,
        EvidencePack("business-integrity"),
    )
    return template_doc, template_spec, records


def test_user_lists_accept_newlines_and_multiple_delimiters() -> None:
    assert parse_user_list("姓名\n入职日期；部门,姓名") == ["姓名", "入职日期", "部门"]


def test_manual_extraction_field_files_use_comma_protocol() -> None:
    for field_file in sorted(MANUAL_EXTRACTIONS.glob("*/提取字段.txt")):
        content = field_file.read_text(encoding="utf-8").strip()
        assert "\n" not in content, field_file
        assert "," in content, field_file
        assert parse_user_list(content) == [field.strip() for field in content.split(",")]


def test_manual_sales_source_builds_values_and_formulas_without_llm() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        template_doc, template_spec, records = _build_records("02-进阶-销售汇总", Path(tmp))
        assert [record.values["区域"] for record in records] == ["华东区", "华南区", "华北区"]
        assert records[0].values["第一季度（万元）"] == 128.5
        assert records[0].values["第二季度（万元）"] == 146.2

        result = XlsxWriter().write(template_doc, template_spec, records)
        worksheet = load_workbook(result.output_path, data_only=False).active
        assert worksheet["D2"].value == "=B2+C2"
        assert worksheet["E2"].value == "=(C2-B2)/B2"


def test_manual_project_and_supplier_sources_apply_business_keys() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        _, _, projects = _build_records("03-困难-项目台账", Path(tmp) / "projects")
        by_project = {record.values["项目编号"]: record.values for record in projects}
        assert by_project["QN-003"]["风险等级"] == "高"
        assert "评审会结论" in by_project["QN-003"]["备注"]
        assert by_project["QN-004"]["预算（万元）"] is None
        assert by_project["QN-004"]["备注"] == "待确认"

    with tempfile.TemporaryDirectory() as tmp:
        _, _, suppliers = _build_records("04-挑战-供应商评分", Path(tmp))
        assert len(suppliers) == 4
        first = next(record for record in suppliers if record.values["信用代码"] == "91330101A001")
        assert first.values["供应商"] == "云杉计算有限公司"
        assert first.values["采用记录"] == "终审"
        assert first.values["冲突说明"]


def test_manual_dashboard_deduplicates_vouchers_and_traces_exceptions() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        template_doc, template_spec, records = _build_records("05-极限-经营驾驶舱", work)
        sheet_names = {
            table.target_table_id: table.logical_name for table in template_spec.target_tables
        }
        grouped = {}
        for record in records:
            grouped.setdefault(sheet_names[record.target_table_id], []).append(record)

        regions = {record.values["区域"]: record.values for record in grouped["区域经营"]}
        # 420 + 460 + (510 - 18) = 1372.  This assertion intentionally follows
        # the source ledger arithmetic instead of the stale 1312 fixture value.
        assert regions["华东"]["净收入（万元）"] == 1372
        assert regions["华南"]["成本（万元）"] == 718
        assert len(grouped["月度明细"]) == 9
        assert len(grouped["异常说明"]) == 3

        result = XlsxWriter().write(template_doc, template_spec, records)
        workbook = load_workbook(result.output_path, data_only=False)
        assert workbook["区域经营"]["D2"].value == "=B2-C2"
        assert workbook["区域经营"]["E2"].value == "=D2/B2"
        assert workbook["区域经营"]["G2"].value == "=B2/F2"


@pytest.mark.parametrize(
    "case_name",
    [
        "01-入门-报名名单",
        "02-进阶-销售汇总",
        "03-困难-项目台账",
        "04-挑战-供应商评分",
        "05-极限-经营驾驶舱",
    ],
)
def test_manual_table_fixtures_match_expected_without_llm(case_name: str) -> None:
    """Keep the deterministic business path aligned with every manual fixture."""

    with tempfile.TemporaryDirectory() as tmp:
        template_doc, template_spec, records = _build_records(case_name, Path(tmp))
        result = XlsxWriter().write(template_doc, template_spec, records)
        mismatches = _compare_workbook(
            Path(result.output_path),
            MANUAL_TABLES / case_name / "期望结果.xlsx",
        )
        assert not mismatches, json.dumps(mismatches, ensure_ascii=False, indent=2, default=str)


@pytest.mark.parametrize("case_name", ["01-入门-报名名单", "03-困难-项目台账"])
def test_reported_table_failures_match_expected_through_full_pipeline(case_name: str) -> None:
    """Exercise the same agent/writer path used by Web tasks without model variance."""

    case_dir = MANUAL_TABLES / case_name
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        for source_path in case_dir.iterdir():
            if source_path.is_file() and source_path.name != "期望结果.xlsx":
                shutil.copy2(source_path, work / source_path.name)
        orchestrator = build_orchestrator(AppConfig(
            enable_agent_runtime=True,
            agent_runtime_backend="local",
            enable_llm_skill_execution=False,
        ))
        result = orchestrator.run(discover_assets(work))
        mismatches = _compare_workbook(
            Path(result.fill_result.output_path),
            case_dir / "期望结果.xlsx",
        )
        assert not mismatches, json.dumps(mismatches, ensure_ascii=False, indent=2, default=str)


def test_document_aware_plan_executes_manual_structures_without_fake_success() -> None:
    case_dir = MANUAL_DOCUMENTS / "01-入门-会议纪要"
    document = Document(case_dir / "原始文档.docx")
    command = (case_dir / "编辑要求.txt").read_text(encoding="utf-8")
    plan = merge_rule_plans(build_rule_based_plan(command), _build_document_aware_plan(command, document))
    results = [_execute_action(document, action)[1] for action in plan.actions]

    assert results and all(result > 0 for result in results)
    assert document.paragraphs[1].text == "第 29 周产品例会纪要"
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["负责人", "事项", "截止日期"],
        ["林晓宇", "准备演示数据", "2026-07-23"],
        ["陈嘉", "修复移动端布局", "2026-07-22"],
    ]
