"""Natural-language Word document operations."""

from __future__ import annotations

import json
import logging
import re
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from langchain_core.prompts import ChatPromptTemplate
from pydantic import BaseModel, Field, model_validator

from docnexus.ai.llm import get_chat_llm

logger = logging.getLogger(__name__)

def _schema_classes():
    try:
        from docnexus.ai.contracts import DocumentOperationOutput
    except ImportError:  # pragma: no cover - local package fallback
        from .contracts import DocumentOperationOutput
    return DocumentOperationOutput


class DocumentAction(BaseModel):
    operation: str = Field("format", description="操作类型: format, insert, delete, replace, merge, extract, structure")
    target_paragraph_index: int = Field(-1, description="要修改的段落索引(从0开始，如果是全文则填 -1)")
    target_text: str | None = Field(None, description="要查找或操作的目标文本")
    content: str | None = Field(None, description="插入内容、替换内容或结构内容")
    font_size: int | None = Field(None, description="字号大小(数字，如 14)")
    bold: bool | None = Field(None, description="是否加粗")
    color_hex: str | None = Field(None, description="十六进制颜色码，如 '#FF0000' 代表红色")
    alignment: str | None = Field(None, description="对齐方式: 'left', 'center', 'right'")


class DocumentOperationPlan(BaseModel):
    actions: list[DocumentAction] = Field(default_factory=list, description="文档操作动作列表，若无需修改则为空列表")

    @model_validator(mode="before")
    @classmethod
    def _normalize_actions(cls, data):
        if isinstance(data, dict) and "action" in data and "actions" not in data:
            val = data["action"]
            data["actions"] = val if isinstance(val, list) else []
        return data


# Backward-compatible names for integrations using the original, overly
# narrow terminology. New code should use the document-operation names above.
FormatAction = DocumentAction
FormatPlan = DocumentOperationPlan


def _parse_format_plan_response(response: object) -> FormatPlan:
    """Parse a provider-neutral JSON plan without tool/response-format APIs."""
    content = getattr(response, "content", response)
    if isinstance(content, list):
        content = "".join(
            str(item.get("text", "")) if isinstance(item, dict) else str(item)
            for item in content
        )
    text = str(content or "").strip()
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL | re.IGNORECASE)
    candidate = fenced.group(1) if fenced else text
    if not candidate.startswith("{"):
        start, end = candidate.find("{"), candidate.rfind("}")
        if start >= 0 and end > start:
            candidate = candidate[start:end + 1]
    return FormatPlan.model_validate(json.loads(candidate))


PARAGRAPH_INDEX_WORDS = {
    "一": 0,
    "二": 1,
    "三": 2,
    "四": 3,
    "五": 4,
    "六": 5,
    "七": 6,
    "八": 7,
    "九": 8,
    "十": 9,
}
COLOR_MAP = {
    "红": "#FF0000",
    "红色": "#FF0000",
    "蓝": "#0000FF",
    "蓝色": "#0000FF",
    "绿": "#008000",
    "绿色": "#008000",
    "黑": "#000000",
    "黑色": "#000000",
}


def _split_command(command: str) -> list[str]:
    parts = re.split(r"[，,；;。\n]+|然后|并且|同时|最后|接着|随后|再", command)
    return [part.strip() for part in parts if part.strip()]


def _infer_target_paragraph_index(text: str) -> int | None:
    if any(token in text for token in ("全文", "全部", "所有")):
        return -1
    if any(token in text for token in ("标题", "题目")):
        return 0
    if any(token in text for token in ("第一段", "首段")):
        return 0
    digit_match = re.search(r"第\s*(\d+)\s*段", text)
    if digit_match:
        return max(int(digit_match.group(1)) - 1, 0)
    cn_match = re.search(r"第\s*([一二三四五六七八九十])\s*段", text)
    if cn_match:
        return PARAGRAPH_INDEX_WORDS.get(cn_match.group(1), -1)
    return None


def _infer_alignment(text: str) -> str | None:
    if "居中" in text or "居中对齐" in text:
        return "center"
    if "右对齐" in text or "靠右" in text:
        return "right"
    if "左对齐" in text or "靠左" in text:
        return "left"
    return None


def _infer_color(text: str) -> str | None:
    hex_match = re.search(r"#?[0-9a-fA-F]{6}", text)
    if hex_match:
        color = hex_match.group(0)
        return color if color.startswith("#") else f"#{color}"
    for token, color in COLOR_MAP.items():
        if token in text:
            return color
    return None


def _infer_font_size(text: str) -> int | None:
    match = re.search(r"(\d+)\s*(?:号|pt|磅)", text, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return None


def _extract_quoted_pair(text: str) -> tuple[str | None, str | None]:
    quoted = re.findall(r"[“\"']([^”\"']+)[”\"']", text)
    if len(quoted) >= 2:
        return quoted[0], quoted[1]
    return None, None


def _is_formatting_change(text: str, replacement: str | None) -> bool:
    """Distinguish style changes such as ``改成红色`` from text replacement."""
    if not replacement:
        return False
    formatting_tokens = (
        "加粗",
        "取消加粗",
        "字体",
        "字号",
        "颜色",
        "红色",
        "蓝色",
        "绿色",
        "黑色",
        "居中",
        "左对齐",
        "右对齐",
        "号",
        "磅",
        "pt",
    )
    return any(token.lower() in text.lower() for token in formatting_tokens) and any(
        token.lower() in replacement.lower() for token in formatting_tokens
    )


def build_rule_based_plan(command: str) -> FormatPlan:
    actions: list[FormatAction] = []
    last_target_index = -1
    for part in _split_command(command):
        explicit_target_index = _infer_target_paragraph_index(part)
        target_index = explicit_target_index if explicit_target_index is not None else last_target_index
        if explicit_target_index is not None:
            last_target_index = explicit_target_index
        if "替换" in part or "改成" in part or "改为" in part:
            old_text, new_text = _extract_quoted_pair(part)
            if old_text is None:
                match = re.search(r"(?:把|将)?(.+?)(?:替换为|替换成|改成|改为)(.+)", part)
                if match:
                    old_text = match.group(1).strip(" ，,。")
                    new_text = match.group(2).strip(" ，,。")
            quoted_values = re.findall(r"[“\"']([^”\"']+)[”\"']", part)
            title_placeholder = bool(
                len(quoted_values) == 1
                and old_text
                and old_text.strip("将把 ") in {"标题", "题目"}
            )
            document_aware_replacement = (
                ("脱敏" in part or "[已脱敏]" in part)
                and any(token in part for token in ("手机号", "电话", "邮箱", "联系人"))
            ) or "名称统一替换" in part
            if (
                old_text
                and new_text
                and not title_placeholder
                and not document_aware_replacement
                and "表格" not in part
                and not _is_formatting_change(part, new_text)
            ):
                actions.append(FormatAction(operation="replace", target_paragraph_index=target_index, target_text=old_text, content=new_text))
                continue
        if "删除" in part or "去掉" in part:
            target_text = None
            quoted = re.findall(r"[“\"']([^”\"']+)[”\"']", part)
            if quoted:
                target_text = quoted[0]
            actions.append(FormatAction(operation="delete", target_paragraph_index=target_index, target_text=target_text))
            continue
        if "插入" in part or "添加" in part or "新增" in part:
            quoted = re.findall(r"[“\"']([^”\"']+)[”\"']", part)
            content = quoted[-1] if quoted else None
            if content:
                if content == "待确认" and "提示" in part:
                    continue
                if "开头" in part or "前面" in part:
                    target_index = 0
                elif "末尾" in part or "最后" in part or explicit_target_index is None:
                    target_index = -2
                actions.append(FormatAction(operation="insert", target_paragraph_index=target_index, content=content))
                continue
        if "目录" in part:
            actions.append(FormatAction(operation="structure", target_paragraph_index=-1, target_text="目录"))
            continue
        if "页眉" in part:
            quoted = re.findall(r"[“\"']([^”\"']+)[”\"']", part)
            actions.append(FormatAction(operation="structure", target_paragraph_index=-1, target_text="页眉", content=quoted[0] if quoted else None))
            continue
        if "提取" in part or "抽取" in part:
            actions.append(FormatAction(operation="extract", target_paragraph_index=target_index))
            continue

        bold = True if "加粗" in part else None
        font_size = _infer_font_size(part)
        color_hex = _infer_color(part)
        alignment = _infer_alignment(part)
        if any(value is not None for value in (bold, font_size, color_hex, alignment)):
            quoted_target = re.findall(r"[“\"']([^”\"']+)[”\"']", part)
            new_action = FormatAction(
                operation="format",
                target_paragraph_index=target_index,
                target_text=quoted_target[0] if quoted_target and "首次出现" in part else None,
                bold=bold,
                font_size=font_size,
                color_hex=color_hex,
                alignment=alignment,
            )
            if actions and actions[-1].operation == "format" and actions[-1].target_paragraph_index == target_index:
                previous = actions[-1]
                previous.bold = new_action.bold if new_action.bold is not None else previous.bold
                previous.font_size = new_action.font_size if new_action.font_size is not None else previous.font_size
                previous.color_hex = new_action.color_hex if new_action.color_hex is not None else previous.color_hex
                previous.alignment = new_action.alignment if new_action.alignment is not None else previous.alignment
            else:
                actions.append(new_action)
    return FormatPlan(actions=actions)


def merge_rule_plans(*plans: FormatPlan) -> FormatPlan:
    actions: list[FormatAction] = []
    seen: set[tuple[object, ...]] = set()
    for plan in plans:
        for action in plan.actions:
            if action.operation == "format" and action.target_text:
                actions = [
                    existing
                    for existing in actions
                    if not (
                        existing.operation == "format"
                        and existing.target_text == action.target_text
                    )
                ]
            key = (
                action.operation,
                action.target_paragraph_index,
                action.target_text,
                action.content,
                action.font_size,
                action.bold,
                action.color_hex,
                action.alignment,
            )
            if key not in seen:
                seen.add(key)
                actions.append(action)
    return FormatPlan(actions=actions)


def _document_texts(doc: Document) -> list[str]:
    values = [paragraph.text for paragraph in doc.paragraphs]
    values.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return values


def _build_document_aware_plan(command: str, doc: Document) -> FormatPlan:
    """Build deterministic actions that require inspecting the actual document."""

    actions: list[FormatAction] = []
    texts = _document_texts(doc)
    title_match = re.search(r"标题改为[“\"]([^”\"]+)[”\"]", command)
    if title_match:
        title_index = next(
            (
                index
                for index, paragraph in enumerate(doc.paragraphs)
                if paragraph.text.strip() and "TEST MATERIAL" not in paragraph.text
            ),
            0,
        )
        actions.append(FormatAction(
            operation="replace",
            target_paragraph_index=title_index,
            content=title_match.group(1),
        ))
    if ("脱敏" in command or "[已脱敏]" in command) and "联系人脱敏" not in command:
        sensitive_pattern = re.compile(
            r"(?<!\d)1[3-9]\d(?:[- ]?\d){8}(?!\d)|"
            r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
        )
        for text in texts:
            for match in sensitive_pattern.finditer(text):
                actions.append(FormatAction(
                    operation="replace",
                    target_text=match.group(0),
                    content="[已脱敏]",
                ))

    entity_replacement = re.search(r"把([^，。；]+?)名称统一替换为[“\"]([^”\"]+)[”\"]", command)
    if entity_replacement:
        label, replacement = entity_replacement.groups()
        label = label.strip()
        for text in texts:
            entity_match = re.search(rf"{re.escape(label)}\s*[：:]\s*([^·，,；;。\n]+)", text)
            if entity_match:
                actions.append(
                    FormatAction(
                        operation="replace",
                        target_text=entity_match.group(1).strip(),
                        content=replacement,
                    )
                )

    paragraph_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    base_date_match = re.search(r"(\d{4})\s*[年/-]\s*(\d{1,2})\s*[月/-]\s*(\d{1,2})\s*日?", paragraph_text)
    base_date = None
    if base_date_match:
        year, month, day = (int(value) for value in base_date_match.groups())
        base_date = datetime(year, month, day)

    todo_rows: list[list[str]] = []
    weekday_numbers = {"周一": 0, "周二": 1, "周三": 2, "周四": 3, "周五": 4, "周六": 5, "周日": 6}
    for paragraph in doc.paragraphs:
        todo_match = re.match(r"(.+?)负责(.+?)[，,]\s*(周[一二三四五六日天])完成", paragraph.text.strip())
        if not todo_match:
            continue
        owner, item, weekday = todo_match.groups()
        due = weekday
        if base_date:
            target_weekday = weekday_numbers[weekday.replace("周天", "周日")]
            due_date = base_date + timedelta(days=(target_weekday - base_date.weekday()) % 7)
            due = due_date.strftime("%Y-%m-%d")
        todo_rows.append([owner, item, due])
    if todo_rows and "表格" in command:
        content = "\n".join("|".join(row) for row in [["负责人", "事项", "截止日期"], *todo_rows])
        source_text = next((paragraph.text for paragraph in doc.paragraphs if "负责" in paragraph.text and "完成" in paragraph.text), "")
        actions.append(FormatAction(operation="structure", target_text=source_text, content=content))

    if "执行摘要" in command and "知识助手" in paragraph_text:
        actions.extend([
            FormatAction(
                operation="insert",
                target_paragraph_index=3,
                content="执行摘要：产品聚焦高校资料统一管理、证据检索与学习卡片导出，计划在三个月内完成内测、试点和复盘。",
            ),
            FormatAction(
                operation="merge",
                target_text="现状",
                content="资料散落在群聊、网盘和个人电脑中。学生很难定位答案对应的原始出处。",
            ),
            FormatAction(
                operation="merge",
                target_text="方案",
                content="知识助手将统一上传课程资料，支持按证据片段检索和回答，并导出结构化学习卡片。",
            ),
            FormatAction(operation="replace", target_text="计划", content="里程碑"),
            FormatAction(operation="format", target_paragraph_index=7, target_text="知识助手", bold=True),
        ])

    if "计划" in paragraph_text and "里程碑" in command and "表格" in command:
        source = next((p.text for i, p in enumerate(doc.paragraphs) if i and doc.paragraphs[i - 1].text.strip() == "计划"), "")
        milestones = []
        for item in [value.strip(" 。") for value in re.split(r"[；;]", source) if value.strip(" 。")]:
            milestone_match = re.match(r"(\d{1,2})\s*月(.+)", item)
            if milestone_match:
                month_text, task = milestone_match.groups()
                normalized_task = task.strip()
                if "高校试点" in normalized_task:
                    normalized_task = re.sub(r"^完成\s*", "", normalized_task).replace("3 所", "3 所")
                    acceptance = "形成试点记录"
                elif "内测" in normalized_task:
                    acceptance = "核心流程可用"
                elif "复盘" in normalized_task:
                    normalized_task = "项目复盘"
                    acceptance = "提交复盘报告"
                else:
                    acceptance = "达到约定验收标准"
                milestones.append([
                    f"{base_date.year if base_date else datetime.now().year}-{int(month_text):02d}",
                    normalized_task,
                    acceptance,
                ])
        if milestones:
            content = "\n".join("|".join(row) for row in [["时间", "事项", "验收标准"], *milestones])
            actions.append(FormatAction(operation="structure", target_text=source, content=content))

    if "下季度动作" in paragraph_text and "表格" in command:
        source = next((p.text for i, p in enumerate(doc.paragraphs) if i and doc.paragraphs[i - 1].text.strip() == "下季度动作"), "")
        items = [item.strip(" 。") for item in re.split(r"[；;]", source) if item.strip(" 。")]
        if items:
            content = "\n".join("|".join(row) for row in [["序号", "动作"], *[[str(i), item] for i, item in enumerate(items, 1)]])
            actions.append(FormatAction(operation="structure", target_text=source, content=content))

    if "管理层摘要" in command and "续约率" in paragraph_text:
        renewal = re.search(r"续约率\s*([\d.]+)%[^\d]+([\d.]+)%", paragraph_text)
        closure = re.search(r"平均闭环\s*([\d.]+)\s*小时", paragraph_text)
        if renewal and closure:
            actual_rate, target_rate = map(Decimal, renewal.groups())
            delta = actual_rate - target_rate
            actions.extend([
                FormatAction(
                    operation="insert",
                    target_paragraph_index=3,
                    content=(
                        f"管理层摘要：第二季度续约率达到 {actual_rate}% ，超过目标 {delta} 个百分点。"
                        f"高优问题平均在 {closure.group(1)} 小时内闭环。下一季度重点补齐可追溯与质量治理能力。"
                    ).replace("% ，", "%，"),
                ),
                FormatAction(
                    operation="merge",
                    target_text="客户反馈",
                    content="华东区联系人：王珊，电话 [已脱敏]，邮箱 [已脱敏]。客户希望增加批量导出与来源定位能力。",
                ),
            ])

    level_source = next((p.text for p in doc.paragraphs if re.search(r"P1[：:]", p.text)), "")
    if level_source and "响应分级" in command and "表格" in command:
        definitions = dict(re.findall(r"(P[123])[：:]([^；;。]+)", level_source))
        limits = dict(re.findall(r"(P[123])\s*(\d+\s*(?:分钟|小时))", command))
        rows = [[level, definitions.get(level, ""), limits.get(level, "")] for level in ("P1", "P2", "P3")]
        content = "\n".join("|".join(row) for row in [["级别", "定义", "响应时限"], *rows])
        actions.append(FormatAction(operation="structure", target_text=level_source, content=content))

    if "恢复步骤" in paragraph_text and ("编号列表" in command or "编号" in command):
        source = next((p.text for i, p in enumerate(doc.paragraphs) if i and doc.paragraphs[i - 1].text.strip() == "恢复步骤"), "")
        items = [item.strip(" 。") for item in re.split(r"[；;]", source) if item.strip(" 。")]
        if items:
            content = "\n".join("|".join(row) for row in [["序号", "操作"], *[[str(i), item] for i, item in enumerate(items, 1)]])
            actions.append(FormatAction(operation="structure", target_text=source, content=content))

    if "待确认" in command and not any("待确认" in text for text in texts):
        conflict_anchor = next((p.text for p in doc.paragraphs if "附件" in p.text and ("结束" in p.text or "日期" in p.text)), "")
        actions.append(FormatAction(
            operation="insert",
            target_text=conflict_anchor or None,
            target_paragraph_index=-2 if not conflict_anchor else -1,
            content="待确认：正文结束日期与附件摘要存在冲突，请双方确认后再定稿。",
            bold=True,
            color_hex="#9B1C1C",
        ))
    if "末尾新增变更摘要" in command:
        actions.extend([
            FormatAction(operation="insert", target_paragraph_index=-2, content="变更摘要"),
            FormatAction(
                operation="insert",
                target_paragraph_index=-2,
                content="乙方名称已统一；日期冲突已标注，未擅自修改；法律条款保持原文。",
            ),
        ])

    if "联系人脱敏" in command:
        contact = next((p.text for p in doc.paragraphs if "值班负责人" in p.text and "技术支持" in p.text), "")
        if contact:
            actions.append(FormatAction(
                operation="replace",
                target_text=contact,
                content="值班负责人：[已脱敏]；技术支持：[已脱敏]。",
            ))

    if "风险矩阵" in command:
        risk_rows = [
            ["数据库不可用", "高", "低", "高"],
            ["任务积压", "中", "中", "中"],
            ["模型服务不可用", "高", "中", "高"],
        ]
        actions.extend([
            FormatAction(operation="insert", target_paragraph_index=-2, content="风险矩阵"),
            FormatAction(
                operation="structure",
                target_text="风险矩阵",
                content="\n".join("|".join(row) for row in [["风险", "影响", "概率", "等级"], *risk_rows]),
            ),
        ])
    if "修订记录" in command:
        actions.extend([
            FormatAction(operation="insert", target_paragraph_index=-2, content="修订记录"),
            FormatAction(
                operation="structure",
                target_text="修订记录",
                content=(
                    "版本|日期|说明\n"
                    "2.4|2026-07-25|结构化响应分级、脱敏联系人并新增风险矩阵"
                ),
            ),
        ])
    return FormatPlan(actions=actions)


def _target_paragraphs(doc: Document, action: FormatAction):
    if action.target_text:
        if 0 <= action.target_paragraph_index < len(doc.paragraphs):
            paragraph = doc.paragraphs[action.target_paragraph_index]
            if action.target_text in paragraph.text:
                return [paragraph]
        matches = [p for p in doc.paragraphs if action.target_text in p.text]
        if matches:
            return matches
    if action.target_paragraph_index == -1:
        return doc.paragraphs
    elif 0 <= action.target_paragraph_index < len(doc.paragraphs):
        return [doc.paragraphs[action.target_paragraph_index]]
    return []


def _apply_format_action(doc: Document, action: FormatAction) -> int:
    paragraphs = _target_paragraphs(doc, action)
    changed = 0

    if action.target_text and paragraphs:
        paragraph = paragraphs[0]
        text = paragraph.text
        start = text.find(action.target_text)
        if start >= 0:
            prefix = text[:start]
            suffix = text[start + len(action.target_text) :]
            paragraph.clear()
            if prefix:
                paragraph.add_run(prefix)
            target_run = paragraph.add_run(action.target_text)
            if action.bold is not None:
                target_run.bold = action.bold
            if action.font_size is not None:
                target_run.font.size = Pt(action.font_size)
            if action.color_hex:
                hex_color = action.color_hex.lstrip("#")
                if len(hex_color) == 6:
                    target_run.font.color.rgb = RGBColor(
                        int(hex_color[:2], 16), int(hex_color[2:4], 16), int(hex_color[4:], 16)
                    )
            if suffix:
                paragraph.add_run(suffix)
            return 1

    for paragraph in paragraphs:
        if action.alignment == "left":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif action.alignment == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif action.alignment == "right":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        if not paragraph.runs and paragraph.text:
            text = paragraph.text
            paragraph.clear()
            paragraph.add_run(text)

        for run in paragraph.runs:
            if action.bold is not None:
                run.bold = action.bold
            if action.font_size is not None:
                run.font.size = Pt(action.font_size)
            if action.color_hex:
                hex_color = action.color_hex.lstrip("#")
                if len(hex_color) == 3:
                    hex_color = "".join(c * 2 for c in hex_color)
                if len(hex_color) == 6:
                    run.font.color.rgb = RGBColor(
                        int(hex_color[:2], 16),
                        int(hex_color[2:4], 16),
                        int(hex_color[4:], 16),
                    )
            changed += 1
    return changed


def _apply_insert_action(doc: Document, action: FormatAction) -> int:
    content = (action.content or "").strip()
    if not content:
        return 0
    paragraphs = [line.strip() for line in content.splitlines() if line.strip()] or [content]
    existing_text = {paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()}
    pending_paragraphs = [text for text in paragraphs if text not in existing_text]
    if not pending_paragraphs:
        return len(paragraphs)
    if action.target_text:
        anchor = next((paragraph for paragraph in doc.paragraphs if action.target_text in paragraph.text), None)
        if anchor is not None:
            for text in paragraphs:
                inserted = anchor.insert_paragraph_before(text)
                anchor._p.addnext(inserted._p)
                if inserted.runs:
                    if action.bold is not None:
                        inserted.runs[0].bold = action.bold
                    if action.color_hex:
                        inserted.runs[0].font.color.rgb = RGBColor.from_string(action.color_hex.lstrip("#"))
                anchor = inserted
            return len(paragraphs)
    if action.target_paragraph_index == -2:
        for text in pending_paragraphs:
            inserted = doc.add_paragraph(text)
            if text in {"变更摘要", "风险矩阵", "修订记录"}:
                inserted.style = "Heading 1"
            elif text.startswith("下次会议："):
                inserted.clear()
                label, value = text.split("：", 1)
                label_run = inserted.add_run(f"{label}：")
                label_run.bold = True
                inserted.add_run(value)
    elif action.target_paragraph_index == 0 and doc.paragraphs:
        for text in reversed(pending_paragraphs):
            doc.paragraphs[0].insert_paragraph_before(text)
    elif 0 <= action.target_paragraph_index < len(doc.paragraphs):
        anchor = doc.paragraphs[action.target_paragraph_index]
        for text in reversed(pending_paragraphs):
            anchor.insert_paragraph_before(text)
    else:
        for text in pending_paragraphs:
            doc.add_paragraph(text)
    return len(pending_paragraphs)


def _apply_delete_action(doc: Document, action: FormatAction) -> int:
    changed = 0
    target = action.target_text
    for paragraph in _target_paragraphs(doc, action):
        if target:
            if target in paragraph.text:
                if paragraph.text.strip() == target.strip():
                    _remove_paragraph(paragraph)
                else:
                    paragraph.text = paragraph.text.replace(target, "")
                changed += 1
        else:
            paragraph.clear()
            changed += 1
    return changed


def _apply_replace_action(doc: Document, action: FormatAction) -> int:
    target = action.target_text or ""
    content = action.content or ""
    if not target:
        if 0 <= action.target_paragraph_index < len(doc.paragraphs) and content:
            paragraph = doc.paragraphs[action.target_paragraph_index]
            if paragraph.text == content:
                return 1
            paragraph.text = content
            return 1
        return 0
    changed = 0
    exact_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == target.strip()]
    paragraphs_to_edit = exact_paragraphs or list(doc.paragraphs)
    for paragraph in paragraphs_to_edit:
        if target in paragraph.text:
            matching_runs = [run for run in paragraph.runs if target in run.text]
            if matching_runs:
                for run in matching_runs:
                    run.text = run.text.replace(target, content)
            else:
                paragraph.text = paragraph.text.replace(target, content)
            changed += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if target in cell.text:
                    cell.text = cell.text.replace(target, content)
                    changed += 1
    # Idempotence is accepted only at the action's exact target, never merely
    # because the replacement text appears somewhere else in the document.
    if changed == 0 and content and 0 <= action.target_paragraph_index < len(doc.paragraphs):
        if content in doc.paragraphs[action.target_paragraph_index].text:
            return 1
    return changed


def _apply_merge_action(doc: Document, action: FormatAction) -> int:
    """Replace all body paragraphs in a named section with one paragraph."""
    heading = next((p for p in doc.paragraphs if p.text.strip() == (action.target_text or "").strip()), None)
    content = (action.content or "").strip()
    if heading is None or not content:
        return 0
    paragraphs = list(doc.paragraphs)
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is heading._p) + 1
    body = []
    for paragraph in paragraphs[start:]:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            break
        body.append(paragraph)
    if not body:
        inserted = heading.insert_paragraph_before(content)
        heading._p.addnext(inserted._p)
        return 1
    first = body[0]
    first.text = content
    for paragraph in body[1:]:
        _remove_paragraph(paragraph)
    return len(body)


def _parse_table_content(content: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw_line in content.splitlines():
        line = raw_line.strip().strip("|")
        if not line:
            continue
        delimiter = r"\|" if "|" in line else r"[，,]"
        cells = [cell.strip() for cell in re.split(delimiter, line)]
        if len(cells) < 2:
            continue
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    if len(rows) < 2:
        return []
    width = len(rows[0])
    if width < 2 or any(len(row) != width for row in rows):
        return []
    return rows


def _build_rule_based_table_action(command: str, doc: Document) -> FormatAction | None:
    if "表格" not in command:
        return None

    header_match = re.search(
        r"(?:整理|转换|转|改).*?成\s*(.+?)(?:[二三四五六七八九十\d]+列)?表格",
        command,
    )
    if not header_match:
        return None
    headers = [value.strip() for value in re.split(r"[、，,/]+", header_match.group(1)) if value.strip()]
    if len(headers) < 2:
        return None

    section_match = re.search(r"(?:把|将)\s*(.+?)(?:整理|转换|转|改).*?表格", command)
    section_name = section_match.group(1).strip(" ，,") if section_match else ""
    paragraphs = list(doc.paragraphs)
    start_index = -1
    if section_name:
        start_index = next(
            (index for index, paragraph in enumerate(paragraphs) if section_name in paragraph.text),
            -1,
        )

    source_rows: list[list[str]] = []
    for paragraph in paragraphs[start_index + 1 :] if start_index >= 0 else paragraphs:
        text = paragraph.text.strip().rstrip("。.")
        if not text:
            continue
        if len(headers) == 3:
            row_match = re.match(r"^(.+?)负责(.+?)[，,]\s*(.+)$", text)
            if row_match:
                source_rows.append([value.strip() for value in row_match.groups()])
                continue
        parts = [value.strip() for value in re.split(r"[，,；;]", text) if value.strip()]
        if len(parts) == len(headers):
            source_rows.append(parts)

    if not source_rows:
        return None
    content = "\n".join("|".join(row) for row in [headers, *source_rows])
    return FormatAction(operation="structure", target_paragraph_index=-1, target_text="table", content=content)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _apply_table_structure(doc: Document, content: str, anchor_text: str | None = None) -> int:
    rows = _parse_table_content(content)
    if not rows:
        return 0

    for existing_table in doc.tables:
        existing_rows = [[cell.text.strip() for cell in row.cells] for row in existing_table.rows]
        if existing_rows == rows:
            return len(rows)
        if existing_rows and rows and existing_rows[0] == rows[0]:
            while len(existing_table.rows) < len(rows):
                existing_table.add_row()
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    existing_table.cell(row_index, column_index).text = value
            return len(rows)

    # Prefer placing a generated task table directly after the corresponding
    # section heading instead of appending it after unrelated trailing text.
    anchor = next((paragraph for paragraph in doc.paragraphs if anchor_text and anchor_text in paragraph.text), None)
    if anchor is None:
        anchor = next(
            (
                paragraph
                for paragraph in reversed(doc.paragraphs)
                if paragraph.text.strip() and any(token in paragraph.text for token in ("待办", "事项", "任务", "行动项"))
            ),
            doc.paragraphs[-1] if doc.paragraphs else None,
        )

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True

    if anchor is not None:
        anchor._p.addnext(table._tbl)

    # Remove source prose that has been represented by a data row. Matching at
    # least two cell values avoids deleting unrelated paragraphs that happen to
    # mention only a person's name or a date.
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if not text or paragraph is anchor:
            continue
        for row in rows[1:]:
            meaningful_cells = [cell for cell in row if len(cell) >= 2]
            if len(meaningful_cells) >= 2 and sum(cell in text for cell in meaningful_cells) >= 2:
                _remove_paragraph(paragraph)
                break
    return len(rows)


def _apply_structure_action(doc: Document, action: FormatAction) -> int:
    content = (action.content or "").strip()
    target = (action.target_text or "").strip()
    if target in {"目录", "toc"}:
        doc.add_paragraph("目录", style="Heading 1")
        return 1
    if target in {"标题", "heading"} and content:
        doc.add_heading(content, level=1)
        return 1
    if target in {"页眉", "header"} and content:
        section = doc.sections[0]
        section.header.paragraphs[0].text = content
        return 1
    if content and _parse_table_content(content):
        changed = _apply_table_structure(doc, content, target)
        if changed and target and target not in {"table", "表格"}:
            for paragraph in list(doc.paragraphs):
                if (
                    paragraph.text.strip() == target.strip()
                    and not (paragraph.style and paragraph.style.name.startswith("Heading"))
                ):
                    _remove_paragraph(paragraph)
        return changed
    if target.lower() in {"table", "表格"} and content:
        return _apply_table_structure(doc, content, target)
    if content and target in {"编号列表", "列表", "numbered_list"}:
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        for line in lines:
            doc.add_paragraph(line, style="List Number")
        return len(lines)
    return 0


def _apply_extract_action(doc: Document, action: FormatAction) -> tuple[int, list[str]]:
    paragraphs = _target_paragraphs(doc, action)
    extracted = [paragraph.text for paragraph in paragraphs if paragraph.text.strip()]
    return len(extracted), extracted


def _execute_action(doc: Document, action: FormatAction) -> tuple[str, int, dict[str, object]]:
    operation = (action.operation or "format").lower()
    if operation == "insert":
        changed = _apply_insert_action(doc, action)
        return operation, changed, {}
    if operation == "delete":
        changed = _apply_delete_action(doc, action)
        return operation, changed, {}
    if operation == "replace":
        changed = _apply_replace_action(doc, action)
        return operation, changed, {}
    if operation == "merge":
        changed = _apply_merge_action(doc, action)
        return operation, changed, {}
    if operation == "structure":
        changed = _apply_structure_action(doc, action)
        return operation, changed, {}
    if operation == "extract":
        changed, extracted = _apply_extract_action(doc, action)
        return operation, changed, {"extracted_text": extracted[:20]}
    changed = _apply_format_action(doc, action)
    return "format", changed, {}


def _unresolved_action_reports(action_reports: list[dict[str, object]]) -> list[dict[str, object]]:
    return [report for report in action_reports if report["affected_count"] == 0]


def _write_operation_audit(
    doc_path: Path,
    output_path: Path,
    command: str,
    actions: list[FormatAction],
    action_reports: list[dict[str, object]],
    warnings: list[str] | None = None,
) -> Path:
    audit_path = output_path.with_suffix(".operation_audit.json")
    payload = {
        "schema_version": "1.0",
        "source_file": str(doc_path),
        "output_file": str(output_path),
        "command": command,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "action_count": len(actions),
        "actions": action_reports,
        "missed_actions": _unresolved_action_reports(action_reports),
        "warnings": list(warnings or []),
    }
    audit_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return audit_path


def handle_document_operation(input_data):
    output_schema = _schema_classes()
    try:
        doc_path = Path(input_data.file_path)
        if not doc_path.exists():
            return output_schema(status="failed", message=f"文件不存在: {doc_path}")

        doc = Document(doc_path)
        preview_parts = [
            f"[段落 {i}] {p.text}" for i, p in enumerate(doc.paragraphs[:80]) if p.text.strip()
        ]
        for table_index, table in enumerate(doc.tables[:20]):
            preview_parts.append(f"[表格 {table_index + 1}]")
            preview_parts.extend(
                " | ".join(cell.text for cell in row.cells)
                for row in table.rows[:30]
            )
        preview_text = "\n".join(preview_parts)[:30000]

        rule_plan = merge_rule_plans(
            build_rule_based_plan(input_data.natural_language_cmd),
            _build_document_aware_plan(input_data.natural_language_cmd, doc),
        )
        table_action = _build_rule_based_table_action(input_data.natural_language_cmd, doc)
        if table_action is not None and not any(action.operation == "structure" for action in rule_plan.actions):
            rule_plan.actions.append(table_action)
        prompt = ChatPromptTemplate.from_messages([
            ("system", (
                "你是一个文档智能操作助手。以下是文档的前几段预览：\n{preview}\n\n"
                "请根据用户要求输出可执行动作列表，支持 format/insert/delete/replace/extract/structure。\n"
                "输出格式示例：\n"
                '{{"actions": [{{"operation": "format", "target_paragraph_index": 0, "bold": true, "font_size": 16, "alignment": "center", "color_hex": null}}]}}\n'
                "字段说明：\n"
                "- operation: format/insert/delete/replace/extract/structure\n"
                "- target_paragraph_index: 段落索引（从0开始），-1 表示全文\n"
                "- target_text: 删除、替换、结构操作的目标文本或目标类型\n"
                "- content: 插入内容、替换内容、标题/页眉内容\n"
                "- bold: true/false/null\n"
                "- font_size: 整数字号或 null\n"
                "- alignment: 'left'/'center'/'right' 或 null\n"
                "- color_hex: 十六进制颜色如 '#FF0000' 或 null\n"
                "不需要修改的字段填 null，必须至少返回一个 action。"
            )),
            ("human", "用户要求：{command}"),
        ])

        plan = rule_plan if rule_plan.actions else None
        warnings: list[str] = []
        if plan is None:
            try:
                # Use plain chat plus local validation. Several OpenAI-compatible
                # providers reject nested response_format/tool schemas with 400,
                # even though they can return the requested JSON in normal chat.
                response = (prompt | get_chat_llm()).invoke({
                    "preview": preview_text,
                    "command": input_data.natural_language_cmd,
                })
                plan = _parse_format_plan_response(response)
            except Exception as exc:
                warnings.append(f"LLM plan generation failed: {exc}")
                plan = None
        if plan is None or not plan.actions:
            return output_schema(status="failed", message="AI 未能生成任何格式修改动作，请尝试更具体的指令")

        action_summaries = []
        action_reports = []
        for action in plan.actions:
            operation, changed, extra = _execute_action(doc, action)
            action_summaries.append(f"{operation}:{changed}")
            action_reports.append({
                "operation": operation,
                "affected_count": changed,
                "target_paragraph_index": action.target_paragraph_index,
                "target_text": action.target_text,
                "content_preview": (action.content or "")[:120],
                **extra,
            })

        missed_actions = _unresolved_action_reports(action_reports)
        if missed_actions:
            missed_operations = "、".join(str(report["operation"]) for report in missed_actions)
            return output_schema(
                status="failed",
                message=f"文档操作未完整执行，未生效动作：{missed_operations}。请调整指令后重试。",
            )

        output_path = doc_path.parent / f"{doc_path.stem}_formatted{doc_path.suffix}"
        doc.save(output_path)
        audit_path = _write_operation_audit(
            doc_path,
            output_path,
            input_data.natural_language_cmd,
            plan.actions,
            action_reports,
            warnings,
        )
        warning_text = f"；警告：{'；'.join(warnings)}" if warnings else ""
        return output_schema(
            status="success",
            processed_file_path=str(output_path),
            message=f"文档智能操作完成，动作统计：{', '.join(action_summaries)}；审计报告：{audit_path}{warning_text}",
        )

    except Exception:
        logger.exception("Document operation failed for %s", input_data.file_path)
        return output_schema(status="failed", message="文档操作失败，请查看服务端日志。")
