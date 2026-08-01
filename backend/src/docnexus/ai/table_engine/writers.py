"""Writers."""

from __future__ import annotations

import logging
import re
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from docnexus.ai.table_engine.core.models import (
    CanonicalDocument,
    CellWriteTrace,
    FillResult,
    StructuredRecord,
    TemplateSpec,
)

logger = logging.getLogger(__name__)


def _ensure_output_path(template_doc: CanonicalDocument) -> Path:
    src_path = Path(template_doc.file.path)
    output_dir = src_path.parent / "outputs"
    output_dir.mkdir(exist_ok=True)
    return output_dir / f"{src_path.stem}-filled{src_path.suffix}"


def _build_fallback_output_path(output_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return output_path.with_name(f"{output_path.stem}-{timestamp}{output_path.suffix}")


def _save_with_fallback(save_func, output_path: Path) -> tuple[Path, list[str]]:
    try:
        save_func(output_path)
        return output_path, []
    except PermissionError:
        fallback_path = _build_fallback_output_path(output_path)
        save_func(fallback_path)
        return fallback_path, [f"Primary output path was locked, wrote fallback file: {fallback_path}"]


def _is_date_field_name(field_name: str) -> bool:
    normalized = "".join(str(field_name).split()).lower()
    return any(token in normalized for token in ("日期", "时间", "date", "time"))


def _normalize_write_value(field_name: str, value: object, data_type: str = "") -> object:
    if value in (None, ""):
        return value
    if data_type.lower() in {"number", "numeric", "float", "int"} and isinstance(value, str):
        text = value.strip().replace(",", "")
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
            number = float(text) if "." in text else int(text)
            return number
    if not _is_date_field_name(field_name):
        return value
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if len(text) >= 10 and text[4:5] in {"-", "/"} and text[7:8] in {"-", "/"}:
        return text[:10].replace("/", "-")
    return value


def _field_role(field_name: str) -> str:
    normalized = re.sub(r"[（(][^）)]*[）)]", "", "".join(str(field_name).lower().split()))
    for role, token in (
        ("budget_execution_rate", "预算执行率"),
        ("budget_achievement_rate", "预算达成率"),
        ("gross_margin_rate", "毛利率"),
        ("qoq_rate", "环比"),
    ):
        if token in normalized:
            return role
    aliases = {
        "month": ("月份", "年月", "month"),
        "date": ("日期", "时间", "date", "time"),
        "region": ("区域", "地区"),
        "country": ("国家/地区", "国家", "country"),
        "project": ("项目编号",),
        "credit": ("信用代码", "统一社会信用代码"),
        "city": ("城市", "城市名"),
        "q1": ("第一季度", "一季度", "q1"),
        "q2": ("第二季度", "二季度", "q2"),
        "income": ("净收入", "收入"),
        "cost": ("成本",),
        "gross": ("毛利",),
        "budget": ("预算",),
        "spent": ("已支出", "支出"),
        "weighted": ("加权总分", "加权得分"),
        "rank": ("排名", "名次"),
    }
    for role, names in aliases.items():
        if any(normalized == name or normalized.startswith(name) for name in names):
            return role
    return normalized


def _identity_fields(fields) -> list[str]:
    by_role = {_field_role(field.field_name): field.field_name for field in fields}
    if "month" in by_role and "region" in by_role:
        return [by_role["month"], by_role["region"]]
    if "date" in by_role:
        for entity_role in ("country", "city", "region"):
            if entity_role in by_role:
                return [by_role[entity_role], by_role["date"]]
    for role in ("credit", "project", "region", "city"):
        if role in by_role:
            return [by_role[role]]
    return [fields[0].field_name] if fields else []


def _record_key(record: StructuredRecord, identity_fields: list[str]) -> tuple[str, ...] | None:
    values = tuple(str(record.values.get(field_name) or "").strip().lower() for field_name in identity_fields)
    return values if values and all(values) else None


def _consolidate_records(records: list[StructuredRecord], fields) -> list[StructuredRecord]:
    identities = _identity_fields(fields)
    ordered: dict[tuple[str, ...], StructuredRecord] = {}
    unkeyed: list[StructuredRecord] = []
    # Record order is a business result (for example, a requested descending
    # ranking). Consolidation must never reshuffle it by confidence.
    for record in records:
        key = _record_key(record, identities)
        if key is None:
            unkeyed.append(deepcopy(record))
            continue
        existing = ordered.get(key)
        if existing is None:
            ordered[key] = deepcopy(record)
            continue
        for field_name, value in record.values.items():
            if value not in (None, "") and (
                existing.values.get(field_name) in (None, "") or record.confidence >= existing.confidence
            ):
                existing.values[field_name] = value
                if record.field_sources.get(field_name):
                    existing.field_sources[field_name] = list(record.field_sources[field_name])
        existing.confidence = max(existing.confidence, record.confidence)
        existing.notes = list(dict.fromkeys([*existing.notes, *record.notes]))
    return [*ordered.values(), *unkeyed]


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _formula_evidence_ids(formula: str, field_names: list[str], record: StructuredRecord) -> list[str]:
    evidence_ids: list[str] = []
    for letters in re.findall(r"\b([A-Z]+)\d+\b", formula.upper()):
        column_index = 0
        for character in letters:
            column_index = column_index * 26 + ord(character) - 64
        if not 1 <= column_index <= len(field_names):
            continue
        for evidence_id in record.field_sources.get(field_names[column_index - 1], []):
            if evidence_id not in evidence_ids:
                evidence_ids.append(evidence_id)
    return evidence_ids


def _derived_formula(
    field_names: list[str],
    field_index: int,
    row_number: int,
    record_count: int,
    record: StructuredRecord,
) -> str | None:
    roles = [_field_role(name) for name in field_names]
    role = roles[field_index]
    columns = {value: _column_letter(index + 1) for index, value in enumerate(roles)}
    row = str(row_number)
    if "合计" in field_names[field_index] and "q1" in columns and "q2" in columns:
        return f"={columns['q1']}{row}+{columns['q2']}{row}"
    if role == "qoq_rate" and "q1" in columns and "q2" in columns:
        return f"=({columns['q2']}{row}-{columns['q1']}{row})/{columns['q1']}{row}"
    if role == "gross" and "率" not in field_names[field_index] and "income" in columns and "cost" in columns:
        return f"={columns['income']}{row}-{columns['cost']}{row}"
    if role == "gross_margin_rate" and "gross" in columns and "income" in columns:
        return f"={columns['gross']}{row}/{columns['income']}{row}"
    if role == "budget_achievement_rate" and "income" in columns and "budget" in columns:
        return f"={columns['income']}{row}/{columns['budget']}{row}"
    if role == "budget_execution_rate" and "spent" in columns and "budget" in columns:
        budget_field = field_names[roles.index("budget")]
        if record.values.get(budget_field) in (None, ""):
            return None
        return f"={columns['spent']}{row}/{columns['budget']}{row}"
    if role == "weighted":
        score_columns = [
            _column_letter(index + 1)
            for index, name in enumerate(field_names)
            if any(token in name for token in ("质量", "交付", "价格", "服务")) and "总" not in name
        ]
        if len(score_columns) == 4:
            weights = ("40%", "30%", "20%", "10%")
            return "=" + "+".join(f"{column}{row}*{weight}" for column, weight in zip(score_columns, weights))
    if role == "rank" and "weighted" in columns and record_count:
        score_column = columns["weighted"]
        end = 1 + record_count
        return f"=RANK({score_column}{row},${score_column}$2:${score_column}${end},0)"
    return None


def _apply_number_format(cell, field_name: str) -> None:
    if any(token in field_name for token in ("手机号", "电话")):
        cell.number_format = "@"
    elif "率" in field_name or "比例" in field_name:
        cell.number_format = "0.0%"
    elif "加权" in field_name or "总分" in field_name:
        cell.number_format = "0.00"
    elif "万元" in field_name:
        cell.number_format = "0.0"


class XlsxWriter:
    supported_doc_types = ("xlsx",)

    def write(
        self,
        template_doc: CanonicalDocument,
        template_spec: TemplateSpec,
        records: list[StructuredRecord],
    ) -> FillResult:
        output_path = _ensure_output_path(template_doc)
        workbook = load_workbook(template_doc.file.path)
        written_cells: list[CellWriteTrace] = []
        records_by_table: dict[str, list[StructuredRecord]] = {}
        for record in records:
            records_by_table.setdefault(record.target_table_id, []).append(record)

        for table_index, target_table in enumerate(template_spec.target_tables):
            if table_index >= len(workbook.worksheets):
                logger.warning(
                    "Target table %s (index %d) has no corresponding worksheet; skipping",
                    target_table.target_table_id, table_index,
                )
                continue
            worksheet = workbook.worksheets[table_index]
            header_row_index = target_table.anchor.row_index if target_table.anchor and target_table.anchor.row_index is not None else 0
            data_start_row = header_row_index + 2
            target_records = _consolidate_records(records_by_table.get(target_table.target_table_id, []), target_table.schema)
            field_names = [field.field_name for field in target_table.schema]
            identity_fields = _identity_fields(target_table.schema)
            identity_columns = [field_names.index(name) + 1 for name in identity_fields if name in field_names]
            used_rows: set[int] = set()
            initial_max_row = worksheet.max_row
            for row_offset, record in enumerate(target_records):
                row_number = 0
                for candidate_row in range(data_start_row, max(initial_max_row, data_start_row - 1) + 1):
                    if candidate_row in used_rows:
                        continue
                    populated = [worksheet.cell(candidate_row, column).value for column in identity_columns]
                    expected = []
                    for column in identity_columns:
                        field_name = field_names[column - 1]
                        expected_value = record.values.get(field_name)
                        if expected_value in (None, "") and _field_role(field_name) in {"序号", "sequence"}:
                            expected_value = row_offset + 1
                        expected.append(expected_value)
                    if populated and all(value not in (None, "") for value in populated) and all(
                        str(left).strip() == str(right).strip() for left, right in zip(populated, expected)
                    ):
                        row_number = candidate_row
                        break
                    # Styled blank template rows are real data slots. Filling
                    # them preserves template formatting and avoids appending
                    # duplicate rows below the intended table body.
                    if all(
                        worksheet.cell(candidate_row, column).value in (None, "")
                        for column in range(1, len(field_names) + 1)
                    ):
                        row_number = candidate_row
                        break
                if not row_number:
                    row_number = max(
                        data_start_row + row_offset,
                        initial_max_row + 1 if initial_max_row >= data_start_row else data_start_row,
                    )
                    while row_number in used_rows:
                        row_number += 1
                if row_number > initial_max_row and data_start_row <= initial_max_row:
                    # New business rows inherit the template's first body-row
                    # presentation.  Writing values into fresh openpyxl cells
                    # otherwise silently drops fills, fonts, borders, alignment,
                    # and number formats after the first row.
                    prototype_row = data_start_row
                    for column in range(1, len(field_names) + 1):
                        source_cell = worksheet.cell(prototype_row, column)
                        target_cell = worksheet.cell(row_number, column)
                        if source_cell.has_style:
                            target_cell._style = deepcopy(source_cell._style)
                    source_dimension = worksheet.row_dimensions[prototype_row]
                    target_dimension = worksheet.row_dimensions[row_number]
                    target_dimension.height = source_dimension.height
                    target_dimension.hidden = source_dimension.hidden
                used_rows.add(row_number)
                for col_index, field in enumerate(target_table.schema, start=1):
                    header_cell = worksheet.cell(row=header_row_index + 1, column=col_index)
                    if header_cell.value in (None, ""):
                        header_cell.value = field.field_name
                    value = _normalize_write_value(
                        field.field_name,
                        record.values.get(field.field_name),
                        field.data_type,
                    )
                    generated_sequence = value in (None, "") and _field_role(field.field_name) in {"序号", "sequence"}
                    if generated_sequence:
                        value = row_offset + 1
                    # Computed columns must remain auditable Excel formulas.
                    # Model-provided scalar answers are snapshots and become
                    # stale when a user edits source cells after delivery.
                    formula = _derived_formula(
                        field_names,
                        col_index - 1,
                        row_number,
                        len(target_records),
                        record,
                    )
                    if formula is not None:
                        value = formula
                    if value in (None, ""):
                        continue
                    target_cell = worksheet.cell(row=row_number, column=col_index, value=value)
                    _apply_number_format(target_cell, field.field_name)
                    evidence_ids = list(record.field_sources.get(field.field_name, []))
                    if formula is not None and not evidence_ids:
                        evidence_ids = _formula_evidence_ids(formula, field_names, record)
                    if generated_sequence and not evidence_ids:
                        evidence_ids = list(dict.fromkeys(
                            evidence_id
                            for sources in record.field_sources.values()
                            for evidence_id in sources
                        ))
                    written_cells.append(
                        CellWriteTrace(
                            target_table_id=target_table.target_table_id,
                            row_index=row_number - 1,
                            col_index=col_index - 1,
                            field_name=field.field_name,
                            value=value,
                            record_id=record.record_id,
                            evidence_ids=evidence_ids,
                        )
                    )

        final_output_path, warnings = _save_with_fallback(workbook.save, output_path)
        return FillResult(
            output_doc_id=template_doc.doc_id,
            output_path=str(final_output_path),
            written_cells=written_cells,
            warnings=warnings,
        )


class DocxTableWriter:
    supported_doc_types = ("docx",)

    def write(
        self,
        template_doc: CanonicalDocument,
        template_spec: TemplateSpec,
        records: list[StructuredRecord],
    ) -> FillResult:
        from docx import Document

        output_path = _ensure_output_path(template_doc)
        document = Document(template_doc.file.path)
        records_by_table: dict[str, list[StructuredRecord]] = {}
        for record in records:
            records_by_table.setdefault(record.target_table_id, []).append(record)

        written_cells: list[CellWriteTrace] = []
        inserted_rows: list[dict[str, object]] = []

        for table_index, target_table in enumerate(template_spec.target_tables):
            if table_index >= len(document.tables):
                logger.warning(
                    "Target table %s (index %d) has no corresponding docx table; skipping",
                    target_table.target_table_id, table_index,
                )
                continue
            docx_table = document.tables[table_index]
            target_records = records_by_table.get(target_table.target_table_id, [])
            if not target_records:
                continue

            header_row_index = target_table.anchor.row_index if target_table.anchor and target_table.anchor.row_index is not None else 0
            data_start_index = header_row_index + 1
            existing_data_rows = max(len(docx_table.rows) - data_start_index, 0)
            while existing_data_rows < len(target_records):
                docx_table.add_row()
                existing_data_rows += 1
                inserted_rows.append({"target_table_id": target_table.target_table_id, "row_index": existing_data_rows})

            for record_offset, record in enumerate(target_records):
                record_index = data_start_index + record_offset
                if record_index >= len(docx_table.rows):
                    logger.warning(
                        "Not enough rows in docx table for target %s; stopping at record %d of %d",
                        target_table.target_table_id, record_index, len(target_records),
                    )
                    break
                row = docx_table.rows[record_index]
                for col_index, field in enumerate(target_table.schema):
                    if col_index >= len(row.cells):
                        continue
                    if header_row_index < len(docx_table.rows) and col_index < len(docx_table.rows[header_row_index].cells):
                        header_cell = docx_table.rows[header_row_index].cells[col_index]
                        if not header_cell.text.strip():
                            header_cell.text = field.field_name
                    value = _normalize_write_value(field.field_name, record.values.get(field.field_name))
                    row.cells[col_index].text = "" if value is None else str(value)
                    written_cells.append(
                        CellWriteTrace(
                            target_table_id=target_table.target_table_id,
                            row_index=record_index,
                            col_index=col_index,
                            field_name=field.field_name,
                            value=value,
                            record_id=record.record_id,
                            evidence_ids=record.field_sources.get(field.field_name, []),
                        )
                    )

        final_output_path, warnings = _save_with_fallback(document.save, output_path)
        return FillResult(
            output_doc_id=template_doc.doc_id,
            output_path=str(final_output_path),
            written_cells=written_cells,
            inserted_rows=inserted_rows,
            warnings=warnings,
        )
