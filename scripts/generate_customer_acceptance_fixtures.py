"""Generate the additional customer-level acceptance fixtures."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "customer-acceptance" / "fixtures" / "20260814-final"


def reset_directory(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True)


def add_styled_paragraph(document: Document, text: str, *, title: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(18 if title else 13)
    if title:
        run.bold = True


def generate_document_cases() -> None:
    case = FIXTURES / "01-document-edit" / "01-project-review"
    case.mkdir(parents=True)
    source = Document()
    add_styled_paragraph(source, "项目阶段复盘", title=True)
    add_styled_paragraph(source, "第一阶段，需求澄清：已完成范围确认。")
    add_styled_paragraph(source, "第二阶段，方案评审：评审通过。")
    source.save(case / "source.docx")
    (case / "command.txt").write_text(
        "将首次出现的“第一阶段，需求澄清”设为深蓝色，再将首次出现的“第二阶段，方案评审”加粗。",
        encoding="utf-8",
    )
    expected = Document(case / "source.docx")
    for paragraph, target, color, bold in (
        (expected.paragraphs[1], "第一阶段，需求澄清", "1F4D78", None),
        (expected.paragraphs[2], "第二阶段，方案评审", None, True),
    ):
        original = paragraph.runs[0]
        suffix = paragraph.text[len(target) :]
        paragraph.clear()
        target_run = paragraph.add_run(target)
        target_run.font.name = original.font.name
        target_run.font.size = original.font.size
        if color:
            target_run.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            target_run.bold = bold
        suffix_run = paragraph.add_run(suffix)
        suffix_run.font.name = original.font.name
        suffix_run.font.size = original.font.size
    expected.save(case / "expected.docx")

    case = FIXTURES / "01-document-edit" / "02-service-weekly"
    case.mkdir(parents=True)
    source = Document()
    add_styled_paragraph(source, "客户服务周报", title=True)
    add_styled_paragraph(source, "本周状态：待确认")
    add_styled_paragraph(source, "下周计划：完成回访并整理问题清单。")
    source.save(case / "source.docx")
    (case / "command.txt").write_text("将标题居中并加粗，然后将“待确认”替换为“已确认”。", encoding="utf-8")
    expected = Document(case / "source.docx")
    expected.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    expected.paragraphs[0].runs[0].bold = True
    expected.paragraphs[1].runs[0].text = "本周状态：已确认"
    expected.save(case / "expected.docx")


def generate_extraction_cases() -> None:
    case = FIXTURES / "02-information-extraction" / "01-contract-change"
    case.mkdir(parents=True)
    (case / "source.txt").write_text(
        """海岸数据中台二期项目合同变更确认单
合同编号：HY-DP-2026-017
项目名称：海岸数据中台二期项目
批准后的含税总金额为人民币 356.8 万元。
补充协议自 2026 年 8 月 15 日起生效。
最终交付地点：江苏省苏州市工业园区星湖街 88 号 A3 栋。
付款条件：合同生效后 10 个工作日内支付 30%，完成中期验收后支付 40%，最终验收通过后支付 30%。
质保期：自最终验收通过之日起 24 个月。
甲方联系人：沈清河，电话 0512-66881234。
项目经理：顾明远。
""",
        encoding="utf-8",
    )
    fields = ["项目名称", "合同编号", "变更后总金额", "生效日期", "交付地点", "付款条件", "质保期", "甲方联系人", "项目经理"]
    (case / "fields.txt").write_text("，".join(fields), encoding="utf-8")
    (case / "expected.json").write_text(
        json.dumps(
            {
                "项目名称": "海岸数据中台二期项目",
                "合同编号": "HY-DP-2026-017",
                "变更后总金额": "3568000.00",
                "生效日期": "2026-08-15",
                "交付地点": "江苏省苏州市工业园区星湖街 88 号 A3 栋",
                "付款条件": "合同生效后 10 个工作日内支付 30%，完成中期验收后支付 40%，最终验收通过后支付 30%",
                "质保期": "24 个月",
                "甲方联系人": "沈清河",
                "项目经理": "顾明远",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    case = FIXTURES / "02-information-extraction" / "02-incident-review"
    case.mkdir(parents=True)
    (case / "source.txt").write_text(
        """支付网关故障复盘
故障发生时间：2026 年 8 月 11 日 14:32。
影响范围：华东区移动端支付请求。
失败请求数：1286 个。
本次故障未发生永久数据丢失，积压任务已全部补偿。
最终根因：任务消费者发布时环境变量名称拼写错误，导致新实例未订阅 production 队列。
服务恢复时间：2026 年 8 月 11 日 15:08。
""",
        encoding="utf-8",
    )
    fields = ["故障发生时间", "影响范围", "失败请求数", "是否发生永久数据丢失", "最终根因", "服务恢复时间"]
    (case / "fields.txt").write_text("，".join(fields), encoding="utf-8")
    (case / "expected.json").write_text(
        json.dumps(
            {
                "故障发生时间": "2026-08-11 14:32",
                "影响范围": "华东区移动端支付请求",
                "失败请求数": 1286,
                "是否发生永久数据丢失": "否",
                "最终根因": "新实例因环境变量名称拼写错误未订阅 production 队列",
                "服务恢复时间": "2026-08-11 15:08",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def generate_table_cases() -> None:
    case = FIXTURES / "03-table-fill" / "01-region-quarter"
    case.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "区域汇总"
    sheet.append(["区域", "第一季度（万元）", "第二季度（万元）", "合计金额（万元）"])
    for row in range(2, 6):
        for column in range(2, 5):
            sheet.cell(row, column).number_format = "0.00"
    workbook.save(case / "template.xlsx")
    (case / "source.txt").write_text(
        "华东区：第一季度 120.5 万元，第二季度 138.2 万元。\n"
        "华南区：第一季度 98.0 万元，第二季度 111.5 万元。\n"
        "华北区：第一季度 105.0 万元，第二季度 100.0 万元。\n"
        "西北区：第一季度 500.0 万元，第二季度 500.0 万元。\n",
        encoding="utf-8",
    )
    (case / "request.txt").write_text("只保留华东区、华南区和华北区，按合计金额从高到低排列。", encoding="utf-8")
    expected = Workbook()
    sheet = expected.active
    sheet.title = "区域汇总"
    sheet.append(["区域", "第一季度（万元）", "第二季度（万元）", "合计金额（万元）"])
    for row, values in enumerate((("华东区", 120.5, 138.2), ("华南区", 98, 111.5), ("华北区", 105, 100)), 2):
        sheet.append([*values, f"=B{row}+C{row}"])
        for column in range(2, 5):
            sheet.cell(row, column).number_format = "0.00"
    expected.save(case / "expected.xlsx")

    case = FIXTURES / "03-table-fill" / "02-supplier-ranking"
    case.mkdir(parents=True)
    template = Workbook()
    sheet = template.active
    sheet.title = "供应商评分"
    sheet.append(["供应商", "质量评分", "交付评分", "价格评分", "服务评分", "加权总分", "排名"])
    for row in range(2, 6):
        sheet.cell(row, 6).number_format = "0.00"
    template.save(case / "template.xlsx")
    source = Workbook()
    sheet = source.active
    sheet.title = "评审数据"
    sheet.append(["供应商", "质量评分", "交付评分", "价格评分", "服务评分", "资格状态"])
    rows = [
        ("启明科技", 88, 92, 84, 90, "有效"),
        ("星桥软件", 90, 82, 92, 88, "有效"),
        ("云帆数据", 99, 99, 99, 99, "取消资格"),
        ("远望系统", 78, 86, 95, 82, "有效"),
    ]
    for row in rows:
        sheet.append(row)
    source.save(case / "source-scores.xlsx")
    (case / "source-status.txt").write_text("云帆数据因资质文件过期，已取消资格，不得进入排名。", encoding="utf-8")
    (case / "request.txt").write_text(
        "仅保留启明科技、星桥软件、云帆数据和远望系统，排除取消资格的记录，按加权总分从高到低排列。",
        encoding="utf-8",
    )
    expected = Workbook()
    sheet = expected.active
    sheet.title = "供应商评分"
    sheet.append(["供应商", "质量评分", "交付评分", "价格评分", "服务评分", "加权总分", "排名"])
    ordered = [rows[0], rows[1], rows[3]]
    for row_index, row in enumerate(ordered, 2):
        sheet.append([*row[:5], f"=B{row_index}*40%+C{row_index}*30%+D{row_index}*20%+E{row_index}*10%", f"=RANK(F{row_index},$F$2:$F$4,0)"])
        sheet.cell(row_index, 6).number_format = "0.00"
    expected.save(case / "expected.xlsx")


def main() -> None:
    reset_directory(FIXTURES)
    generate_document_cases()
    generate_extraction_cases()
    generate_table_cases()
    print(f"generated fixtures at {FIXTURES}")


if __name__ == "__main__":
    main()
