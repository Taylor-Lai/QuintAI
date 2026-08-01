"""Run the repository's 15 manual acceptance fixtures through the live API.

Results are written incrementally to a run-specific directory under
``reports/test-runs``. Raw run artifacts are intentionally not committed.
"""

from __future__ import annotations

import argparse
import json
import os
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "manual"
RUN_ID = datetime.now().strftime("%Y%m%d-%H%M%S")
OUTPUT = ROOT / "reports" / "test-runs" / f"{RUN_ID}-manual"
RESULTS_PATH = OUTPUT / "results.json"
BASE_URL = os.getenv("ACCEPTANCE_BASE_URL", "http://127.0.0.1:8000/api")
EMAIL = os.getenv("ACCEPTANCE_EMAIL", "")
PASSWORD = os.getenv("ACCEPTANCE_PASSWORD", "")
CASE_NAMES: set[str] = set()


def selected_case(path: Path) -> bool:
    return not CASE_NAMES or path.name in CASE_NAMES


def serializable(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    return value


def persist(report: dict[str, Any]) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    RESULTS_PATH.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=serializable),
        encoding="utf-8",
    )


def login(client: httpx.Client) -> None:
    response = client.post("/auth/login", json={"email": EMAIL, "password": PASSWORD})
    response.raise_for_status()
    client.headers["Authorization"] = f"Bearer {response.json()['access_token']}"


def wait_for_task(client: httpx.Client, task_id: str, timeout_seconds: int = 1800) -> dict[str, Any]:
    deadline = time.monotonic() + timeout_seconds
    last_stage = None
    while time.monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}")
        response.raise_for_status()
        task = response.json()
        if task.get("stage") != last_stage:
            print(f"  {task_id[:8]} {task.get('status')} {task.get('progress')}% {task.get('stage')}", flush=True)
            last_stage = task.get("stage")
        if task.get("status") in {"succeeded", "failed", "cancelled"}:
            return task
        time.sleep(1)
    raise TimeoutError(f"task {task_id} exceeded {timeout_seconds}s")


def submit_and_wait(
    client: httpx.Client,
    endpoint: str,
    *,
    data: dict[str, str],
    files: list[tuple[str, tuple[str, Any, str]]],
) -> tuple[dict[str, Any], float]:
    started = time.monotonic()
    for attempt in range(1, 9):
        # httpx consumes upload streams. Rewind them before every retry so a
        # rate-limited request is resubmitted with the original file content.
        for _, (_, stream, _) in files:
            if hasattr(stream, "seek"):
                stream.seek(0)
        response = client.post(endpoint, data=data, files=files, timeout=60)
        if response.status_code != 429:
            response.raise_for_status()
            break
        retry_after = max(int(response.headers.get("Retry-After", "5")), 1)
        print(
            f"  rate limited (attempt {attempt}/8); retrying in {retry_after}s",
            flush=True,
        )
        time.sleep(retry_after)
    else:
        response.raise_for_status()
    task = wait_for_task(client, response.json()["id"])
    return task, round(time.monotonic() - started, 3)


def simple_mime(path: Path) -> str:
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }[path.suffix.lower()]


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u3000", " ").split())


def compare_extraction(actual: dict[str, Any], expected: dict[str, Any]) -> dict[str, Any]:
    mismatches = []
    for field, expected_value in expected.items():
        actual_value = actual.get(field)
        if norm_text(actual_value) != norm_text(expected_value):
            mismatches.append({"field": field, "expected": expected_value, "actual": actual_value})
    unexpected = sorted(key for key in actual if key not in expected and key != "_meta")
    return {
        "passed": not mismatches and not unexpected,
        "matched_fields": len(expected) - len(mismatches),
        "total_fields": len(expected),
        "mismatches": mismatches,
        "unexpected_fields": unexpected,
        "has_meta": isinstance(actual.get("_meta"), dict),
    }


def run_snapshot(run: Any) -> dict[str, Any]:
    color = run.font.color.rgb
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "size": run.font.size.pt if run.font.size else None,
        "color": str(color) if color else None,
        "font": run.font.name,
    }


def document_snapshot(path: Path) -> dict[str, Any]:
    doc = Document(path)
    paragraphs = []
    for paragraph in doc.paragraphs:
        paragraphs.append(
            {
                "text": paragraph.text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
                "runs": [run_snapshot(run) for run in paragraph.runs],
            }
        )
    tables = []
    for table in doc.tables:
        tables.append([[cell.text for cell in row.cells] for row in table.rows])
    headers = [[p.text for p in section.header.paragraphs] for section in doc.sections]
    footers = [[p.text for p in section.footer.paragraphs] for section in doc.sections]
    return {"paragraphs": paragraphs, "tables": tables, "headers": headers, "footers": footers}


def compare_document(actual_path: Path, expected_path: Path) -> dict[str, Any]:
    actual = document_snapshot(actual_path)
    expected = document_snapshot(expected_path)
    actual_text = [item["text"] for item in actual["paragraphs"] if item["text"].strip()]
    expected_text = [item["text"] for item in expected["paragraphs"] if item["text"].strip()]
    differences = []
    if actual_text != expected_text:
        differences.append({"area": "paragraph_text", "expected": expected_text, "actual": actual_text})
    if actual["tables"] != expected["tables"]:
        differences.append({"area": "tables", "expected": expected["tables"], "actual": actual["tables"]})
    if actual["headers"] != expected["headers"]:
        differences.append({"area": "headers", "expected": expected["headers"], "actual": actual["headers"]})
    if actual["footers"] != expected["footers"]:
        differences.append({"area": "footers", "expected": expected["footers"], "actual": actual["footers"]})
    format_equal = actual["paragraphs"] == expected["paragraphs"]
    if not format_equal:
        differences.append({"area": "paragraph_format", "detail": "paragraph/run formatting snapshot differs"})
    return {
        "passed": not differences,
        "file_opens": True,
        "content_equal": actual_text == expected_text and actual["tables"] == expected["tables"],
        "format_equal": format_equal,
        "differences": differences[:8],
    }


def cell_snapshot(cell: Any) -> dict[str, Any]:
    fill_value = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
    fill = str(fill_value) if fill_value is not None else None
    color = cell.font.color
    font_color = None
    if color is not None and color.type == "rgb":
        font_color = str(color.rgb)
    return {
        "value": serializable(cell.value),
        "number_format": cell.number_format,
        "font": {"bold": cell.font.bold, "italic": cell.font.italic, "color": font_color},
        "fill": fill,
        "alignment": {"horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical},
    }


def workbook_snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    snapshot: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        cells = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = cell_snapshot(cell)
        snapshot[sheet.title] = {
            "cells": cells,
            "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
        }
    return snapshot


def compare_workbook(actual_path: Path, expected_path: Path) -> dict[str, Any]:
    actual = workbook_snapshot(actual_path)
    expected = workbook_snapshot(expected_path)
    mismatches = []
    for sheet_name in sorted(set(actual) | set(expected)):
        if sheet_name not in actual or sheet_name not in expected:
            mismatches.append({"sheet": sheet_name, "area": "missing_sheet"})
            continue
        actual_cells = actual[sheet_name]["cells"]
        expected_cells = expected[sheet_name]["cells"]
        for coordinate in sorted(set(actual_cells) | set(expected_cells)):
            a = actual_cells.get(coordinate)
            e = expected_cells.get(coordinate)
            if a != e:
                mismatches.append({"sheet": sheet_name, "cell": coordinate, "expected": e, "actual": a})
        if actual[sheet_name]["merged"] != expected[sheet_name]["merged"]:
            mismatches.append(
                {
                    "sheet": sheet_name,
                    "area": "merged_cells",
                    "expected": expected[sheet_name]["merged"],
                    "actual": actual[sheet_name]["merged"],
                }
            )
    return {
        "passed": not mismatches,
        "file_opens": True,
        "mismatch_count": len(mismatches),
        "mismatches": mismatches[:30],
    }


def task_common(task: dict[str, Any], duration: float) -> dict[str, Any]:
    return {
        "task_id": task.get("id"),
        "status": task.get("status"),
        "duration_seconds": duration,
        "stage": task.get("stage"),
        "error": task.get("error"),
        "quality_report": task.get("quality_report"),
        "evidence_summary": task.get("evidence_summary"),
        "events": task.get("events"),
    }


def extraction_cases(client: httpx.Client, report: dict[str, Any]) -> None:
    root = FIXTURES / "01-信息提取"
    for case_dir in sorted(path for path in root.iterdir() if path.is_dir() and selected_case(path)):
        print(f"[extract] {case_dir.name}", flush=True)
        source = next(case_dir.glob("源材料.*"))
        fields_path = case_dir / "提取字段.txt"
        expected_path = case_dir / "期望结果.json"
        raw_fields = fields_path.read_text(encoding="utf-8").strip()
        with source.open("rb") as source_file:
            task, duration = submit_and_wait(
                client,
                "/doc-extract/upload",
                data={"fields": raw_fields},
                files=[("file", (source.name, source_file, simple_mime(source)))],
            )
        item = task_common(task, duration)
        if task.get("status") == "succeeded":
            expected = json.loads(expected_path.read_text(encoding="utf-8"))
            actual = (task.get("result") or {}).get("extracted_data") or {}
            item["comparison"] = compare_extraction(actual, expected)
        report["scenarios"].append({"category": "information_extraction", "name": case_dir.name, **item})
        persist(report)


def document_cases(client: httpx.Client, report: dict[str, Any]) -> None:
    root = FIXTURES / "02-文档编辑"
    out_dir = OUTPUT / "document_outputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    for case_dir in sorted(path for path in root.iterdir() if path.is_dir() and selected_case(path)):
        print(f"[document] {case_dir.name}", flush=True)
        source = case_dir / "原始文档.docx"
        command = (case_dir / "编辑要求.txt").read_text(encoding="utf-8").strip()
        with source.open("rb") as source_file:
            task, duration = submit_and_wait(
                client,
                "/doc-chat/upload",
                data={"command": command},
                files=[("document", (source.name, source_file, simple_mime(source)))],
            )
        item = task_common(task, duration)
        if task.get("status") == "succeeded" and task.get("has_file"):
            response = client.get(f"/tasks/{task['id']}/download", timeout=60)
            response.raise_for_status()
            output_path = out_dir / f"{case_dir.name}.docx"
            output_path.write_bytes(response.content)
            item["output"] = str(output_path.relative_to(ROOT))
            item["comparison"] = compare_document(output_path, case_dir / "期望结果.docx")
        report["scenarios"].append({"category": "document_editing", "name": case_dir.name, **item})
        persist(report)


def table_cases(client: httpx.Client, report: dict[str, Any]) -> None:
    root = FIXTURES / "03-表格填充"
    out_dir = OUTPUT / "table_outputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    for case_dir in sorted(path for path in root.iterdir() if path.is_dir() and selected_case(path)):
        print(f"[table] {case_dir.name}", flush=True)
        template = case_dir / "空白模板.xlsx"
        request = (case_dir / "用户要求.txt").read_text(encoding="utf-8").strip()
        sources = sorted(
            path
            for path in case_dir.iterdir()
            if path.is_file() and (path.name.startswith("源数据") or path.name.startswith("补充说明"))
        )
        handles = []
        try:
            template_handle = template.open("rb")
            handles.append(template_handle)
            files: list[tuple[str, tuple[str, Any, str]]] = [
                ("template", (template.name, template_handle, simple_mime(template)))
            ]
            for source in sources:
                handle = source.open("rb")
                handles.append(handle)
                files.append(("documents", (source.name, handle, simple_mime(source))))
            task, duration = submit_and_wait(
                client,
                "/table-fill/upload",
                data={"user_request": request},
                files=files,
            )
        finally:
            for handle in handles:
                handle.close()
        item = task_common(task, duration)
        item["source_files"] = [path.name for path in sources]
        if task.get("status") == "succeeded" and task.get("has_file"):
            response = client.get(f"/tasks/{task['id']}/download", timeout=60)
            response.raise_for_status()
            output_path = out_dir / f"{case_dir.name}.xlsx"
            output_path.write_bytes(response.content)
            item["output"] = str(output_path.relative_to(ROOT))
            item["comparison"] = compare_workbook(output_path, case_dir / "期望结果.xlsx")
        report["scenarios"].append({"category": "table_filling", "name": case_dir.name, **item})
        persist(report)


def main() -> None:
    global BASE_URL, CASE_NAMES, EMAIL, OUTPUT, PASSWORD, RESULTS_PATH
    parser = argparse.ArgumentParser(description="Run all 15 tests/manual scenarios through a live API.")
    parser.add_argument("--base-url", default=BASE_URL)
    parser.add_argument("--email", default=EMAIL)
    parser.add_argument("--password", default=PASSWORD)
    parser.add_argument("--run-id", default=RUN_ID)
    parser.add_argument(
        "--category",
        choices=("all", "extraction", "document", "table"),
        default="all",
        help="Run one fixture category instead of all 15 cases.",
    )
    parser.add_argument(
        "--case",
        action="append",
        default=[],
        help="Run only an exact fixture directory name; may be repeated.",
    )
    args = parser.parse_args()
    if not args.email or not args.password:
        parser.error("provide --email/--password or ACCEPTANCE_EMAIL/ACCEPTANCE_PASSWORD")
    BASE_URL = args.base_url.rstrip("/")
    EMAIL = args.email
    PASSWORD = args.password
    CASE_NAMES = set(args.case)
    OUTPUT = ROOT / "reports" / "test-runs" / f"{args.run_id}-manual"
    RESULTS_PATH = OUTPUT / "results.json"
    report: dict[str, Any] = {
        "schema_version": "2.0",
        "run_id": args.run_id,
        "started_at": datetime.now().astimezone().isoformat(),
        "base_url": BASE_URL,
        "notes": [
            "Web registration/login/upload smoke test was run before this API batch.",
            "Fixture expected-result files were used only for comparison, never uploaded.",
        ],
        "scenarios": [],
    }
    persist(report)
    base_host = (urlparse(BASE_URL).hostname or "").lower()
    trust_environment_proxy = base_host not in {"localhost", "127.0.0.1", "::1"}
    with httpx.Client(base_url=BASE_URL, timeout=30, trust_env=trust_environment_proxy) as client:
        login(client)
        if args.category in {"all", "extraction"}:
            extraction_cases(client, report)
        if args.category in {"all", "document"}:
            document_cases(client, report)
        if args.category in {"all", "table"}:
            table_cases(client, report)
    report["completed_at"] = datetime.now().astimezone().isoformat()
    report["summary"] = {
        "total": len(report["scenarios"]),
        "tasks_succeeded": sum(item.get("status") == "succeeded" for item in report["scenarios"]),
        "comparisons_passed": sum((item.get("comparison") or {}).get("passed") is True for item in report["scenarios"]),
    }
    persist(report)
    print(json.dumps(report["summary"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
