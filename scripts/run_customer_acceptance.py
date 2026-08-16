"""Run the six additional customer-level fixtures against a live API."""

from __future__ import annotations

import argparse
import json
import mimetypes
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "customer-acceptance" / "fixtures" / "20260814-final"


def mime(path: Path) -> str:
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def wait_for_task(client: httpx.Client, task_id: str, timeout: int) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    last_snapshot: tuple[str, int, int, int, str] | None = None
    next_heartbeat = 0.0
    while time.monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}")
        response.raise_for_status()
        task = response.json()
        snapshot = (
            str(task.get("status") or "unknown"),
            int(task.get("progress") or 0),
            int(task.get("completed_steps") or 0),
            int(task.get("total_steps") or 1),
            str(task.get("stage") or "等待执行"),
        )
        now = time.monotonic()
        if snapshot != last_snapshot or now >= next_heartbeat:
            status, progress, completed, total, stage = snapshot
            print(f"  [{status}] {progress}% · 步骤 {completed}/{total} · {stage}", flush=True)
            last_snapshot = snapshot
            next_heartbeat = now + 15
        if task.get("status") in {"succeeded", "failed", "cancelled"}:
            return task
        time.sleep(1)
    raise TimeoutError(f"task {task_id} exceeded {timeout}s")


def submit(
    client: httpx.Client,
    endpoint: str,
    data: dict[str, str],
    files: list[tuple[str, tuple[str, Any, str]]],
    timeout: int,
) -> tuple[dict[str, Any], float]:
    started = time.monotonic()
    for attempt in range(8):
        for _, (_, stream, _) in files:
            stream.seek(0)
        response = client.post(endpoint, data=data, files=files, timeout=60)
        if response.status_code != 429:
            response.raise_for_status()
            task = wait_for_task(client, response.json()["id"], timeout)
            return task, round(time.monotonic() - started, 3)
        time.sleep(max(int(response.headers.get("Retry-After", "5")), 1))
    response.raise_for_status()
    raise AssertionError("unreachable")


def normalize_text(value: Any) -> str:
    return " ".join(str(value if value is not None else "").strip("。；;，, ").split())


def compare_extraction(actual: dict[str, Any], expected: dict[str, Any]) -> dict[str, Any]:
    mismatches = []
    for field, expected_value in expected.items():
        actual_value = actual.get(field)
        if normalize_text(actual_value) != normalize_text(expected_value):
            mismatches.append({"field": field, "expected": expected_value, "actual": actual_value})
    unexpected = sorted(field for field in actual if field not in expected and field != "_meta")
    return {"passed": not mismatches and not unexpected, "mismatches": mismatches, "unexpected": unexpected}


def run_snapshot(run: Any) -> tuple[Any, ...]:
    color = run.font.color.rgb
    return (
        run.text,
        run.bold,
        run.italic,
        bool(run.underline) if run.underline is not None else None,
        run.font.size.pt if run.font.size else None,
        str(color) if color else None,
        run.font.name,
    )


def document_snapshot(path: Path) -> dict[str, Any]:
    document = Document(path)
    return {
        "paragraphs": [
            {
                "text": paragraph.text,
                "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
                "runs": [run_snapshot(run) for run in paragraph.runs],
            }
            for paragraph in document.paragraphs
        ],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
    }


def compare_document(actual: Path, expected: Path) -> dict[str, Any]:
    actual_snapshot = document_snapshot(actual)
    expected_snapshot = document_snapshot(expected)
    return {
        "passed": actual_snapshot == expected_snapshot,
        "actual": actual_snapshot if actual_snapshot != expected_snapshot else None,
        "expected": expected_snapshot if actual_snapshot != expected_snapshot else None,
    }


def workbook_snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    snapshot: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        cells: dict[str, Any] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    cells[cell.coordinate] = {"value": cell.value, "number_format": cell.number_format}
        snapshot[sheet.title] = {"cells": cells, "merged": sorted(str(item) for item in sheet.merged_cells.ranges)}
    return snapshot


def compare_workbook(actual: Path, expected: Path) -> dict[str, Any]:
    actual_snapshot = workbook_snapshot(actual)
    expected_snapshot = workbook_snapshot(expected)
    mismatches = []
    for sheet_name, expected_sheet in expected_snapshot.items():
        actual_sheet = actual_snapshot.get(sheet_name)
        if actual_sheet is None:
            mismatches.append({"sheet": sheet_name, "error": "missing sheet"})
            continue
        for coordinate, expected_cell in expected_sheet["cells"].items():
            actual_cell = actual_sheet["cells"].get(coordinate)
            if actual_cell != expected_cell:
                mismatches.append(
                    {"sheet": sheet_name, "cell": coordinate, "expected": expected_cell, "actual": actual_cell}
                )
        if actual_sheet["merged"] != expected_sheet["merged"]:
            mismatches.append(
                {"sheet": sheet_name, "area": "merged", "expected": expected_sheet["merged"], "actual": actual_sheet["merged"]}
            )
    return {"passed": not mismatches, "mismatches": mismatches[:40]}


def task_item(task: dict[str, Any], duration: float) -> dict[str, Any]:
    return {
        "task_id": task.get("id"),
        "status": task.get("status"),
        "duration_seconds": duration,
        "stage": task.get("stage"),
        "error": task.get("error"),
        "quality_report": task.get("quality_report"),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8000/api")
    parser.add_argument("--email", required=True)
    parser.add_argument("--password", required=True)
    parser.add_argument("--username", default="customer-acceptance")
    parser.add_argument("--register", action="store_true")
    parser.add_argument("--run-id", default=datetime.now().strftime("%Y%m%d-%H%M%S"))
    parser.add_argument("--timeout", type=int, default=1800)
    args = parser.parse_args()
    output = ROOT / "reports" / "test-runs" / f"{args.run_id}-customer"
    output.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "run_id": args.run_id,
        "base_url": args.base_url,
        "started_at": datetime.now().astimezone().isoformat(),
        "scenarios": [],
    }

    def persist() -> None:
        (output / "results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    persist()
    with httpx.Client(base_url=args.base_url.rstrip("/"), timeout=30, trust_env=False) as client:
        if args.register:
            response = client.post(
                "/auth/register",
                json={"username": args.username, "email": args.email, "password": args.password},
            )
            if response.status_code not in {200, 201, 400, 409}:
                response.raise_for_status()
        response = client.post("/auth/login", json={"email": args.email, "password": args.password})
        response.raise_for_status()
        client.headers["Authorization"] = f"Bearer {response.json()['access_token']}"

        for case in sorted((FIXTURES / "01-document-edit").iterdir()):
            print(f"[document] {case.name}", flush=True)
            with (case / "source.docx").open("rb") as stream:
                task, duration = submit(
                    client,
                    "/doc-chat/upload",
                    {"command": (case / "command.txt").read_text(encoding="utf-8")},
                    [("document", ("source.docx", stream, mime(case / "source.docx")))],
                    args.timeout,
                )
            item = {"category": "document_editing", "name": case.name, **task_item(task, duration)}
            if task.get("status") == "succeeded" and task.get("has_file"):
                result = client.get(f"/tasks/{task['id']}/download")
                result.raise_for_status()
                actual = output / f"{case.name}.docx"
                actual.write_bytes(result.content)
                item["comparison"] = compare_document(actual, case / "expected.docx")
            report["scenarios"].append(item)
            persist()

        for case in sorted((FIXTURES / "02-information-extraction").iterdir()):
            print(f"[extraction] {case.name}", flush=True)
            source = case / "source.txt"
            with source.open("rb") as stream:
                task, duration = submit(
                    client,
                    "/doc-extract/upload",
                    {"fields": (case / "fields.txt").read_text(encoding="utf-8")},
                    [("file", (source.name, stream, mime(source)))],
                    args.timeout,
                )
            item = {"category": "information_extraction", "name": case.name, **task_item(task, duration)}
            if task.get("status") == "succeeded":
                actual = (task.get("result") or {}).get("extracted_data") or {}
                expected = json.loads((case / "expected.json").read_text(encoding="utf-8"))
                item["comparison"] = compare_extraction(actual, expected)
            report["scenarios"].append(item)
            persist()

        for case in sorted((FIXTURES / "03-table-fill").iterdir()):
            print(f"[table] {case.name}", flush=True)
            paths = [path for path in case.iterdir() if path.name.startswith("source-") or path.name == "source.txt"]
            handles = [(case / "template.xlsx").open("rb"), *(path.open("rb") for path in paths)]
            try:
                files = [("template", ("template.xlsx", handles[0], mime(case / "template.xlsx")))]
                files.extend(("documents", (path.name, handle, mime(path))) for path, handle in zip(paths, handles[1:]))
                task, duration = submit(
                    client,
                    "/table-fill/upload",
                    {"user_request": (case / "request.txt").read_text(encoding="utf-8")},
                    files,
                    args.timeout,
                )
            finally:
                for handle in handles:
                    handle.close()
            item = {"category": "table_filling", "name": case.name, **task_item(task, duration)}
            if task.get("status") == "succeeded" and task.get("has_file"):
                result = client.get(f"/tasks/{task['id']}/download")
                result.raise_for_status()
                actual = output / f"{case.name}.xlsx"
                actual.write_bytes(result.content)
                item["comparison"] = compare_workbook(actual, case / "expected.xlsx")
            report["scenarios"].append(item)
            persist()

    report["completed_at"] = datetime.now().astimezone().isoformat()
    durations = [float(item["duration_seconds"]) for item in report["scenarios"]]
    report["summary"] = {
        "total": len(report["scenarios"]),
        "tasks_succeeded": sum(item.get("status") == "succeeded" for item in report["scenarios"]),
        "comparisons_passed": sum((item.get("comparison") or {}).get("passed") is True for item in report["scenarios"]),
        "duration_seconds": round(sum(durations), 3),
    }
    persist()
    print(json.dumps(report["summary"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
